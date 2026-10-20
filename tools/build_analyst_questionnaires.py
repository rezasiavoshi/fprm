"""Build the analyst questionnaires (.xlsx + .md): Chemical lab Raw-Material (RM) and Finished-Product (FP),
Physical lab (finished-product physical tests and IPC checks) and Micro lab (raw materials, products, monitoring).

Pre-fills each questionnaire from:
  data/QC_Capacity_Study.xlsx  - item lists, tests and column layout (answers paste back 1:1)
  data/Leadtime.xlsx           - recorded RM test times (minutes), used as hints to confirm
  data/rawmaterials.md         - raw-material names as purchased (usp48 export), checked against the study list
  data/usp_tests.csv           - RM compendial tests per material, flattened from USP-NF (tools/usp_tests.py)
  data/usp_micro.csv           - RM microbiological tests per material from USP-NF (tools/usp_tests.py)
  data/Product-Material-Database.xlsx - 'Name Mapping' sheet, source spelling -> study material name
The questionnaires are self-contained: respondent-facing text never names these files or their sheets,
so a respondent needs only the questionnaire. Column order still matches the study for paste-back.
Usage: python3 tools/build_analyst_questionnaires.py [repo_root]
"""
import re, sys, os, csv, collections, difflib, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = sys.argv[1] if len(sys.argv) > 1 else "."
STUDY = os.path.join(ROOT, "data/QC_Capacity_Study.xlsx")
LEAD = os.path.join(ROOT, "data/Leadtime.xlsx")
RAW = os.path.join(ROOT, "data/rawmaterials.md")
DB = os.path.join(ROOT, "data/Product-Material-Database.xlsx")
USP_TESTS = os.path.join(ROOT, "data/usp_tests.csv")   # tools/usp_tests.py
USP_MICRO = os.path.join(ROOT, "data/usp_micro.csv")   # tools/usp_tests.py
OUT = os.path.join(ROOT, "questionnaires")

# Study material name -> Leadtime.xlsx name(s). Curated: automatic fuzzy matching gave false pairs
# (Sucrose/Sucralose, Selenate/Selenite, MCC grades), so every non-exact pair is listed here.
ALIASES = {
    "Apium Graveolens Seed Dry Extract 6:1": ["Apium Graveolens (Celery Seed Ext.)"],
    "Artichoke Dry Extract 25:1": ["Cynara Scolymus (Globe Artichoke Ext.) Dry Conc 25:1"],
    "Ascorbic Acid (Vitamin C)": ["Vitamin C (Ascorbic Acid)"],
    "Ascorbic Acid Coated 97%": ["Vitamin C Coated 97.5% (Ascorbic Acid Coated)"],
    "Azelastine Hydrochloride": ["Azelastine HCL (MSN)"],
    "Beta-Carotene": ["Beta Carotene 10%"],
    "Biotin (Vitamin B7)": ["D-Biotin"],
    "Calcium Carbonate Heavy": ["Calcium Carbonate (Heavy)"],
    "Calcium Pantothenate (Vitamin B5)": ["Vitamin B5 (Calcium Pantothenate)"],
    "Citicoline Sodium": ["Citicoline"],
    "Cyanocobalamin (Vitamin B12)": ["Vitamin B12 (Cyanocobalamine)"],
    "D-Alpha Tocopheryl Acetate 50% (Vitamin E)": ["Vitamin E (D-Alpha Tocopheryl Acetate)"],
    "Dapagliflozin Propanediol Monohydrate": ["Dapagliflozine Propanediol Monohydrate (Tofigh Darou)",
                                              "Dapagliflozine Propanediol Monohydrate (Alborz Darou)"],
    "Devil's Claw Dry Extract 6:1": ["Harpagophytum Procombens (Devils Claw Ext.) Dry Conc 6:1"],
    "Dextromethorphan HBr": ["Dextromethorphan"],
    "Diphenhydramine HCl": ["Diphenhydramine"],
    "Echinacea Purpurea Dry Extract": ["Echinacea"],
    "Epimedium Grandiflorum (Horny Goat Weed) Dry Extract 20:1": ["Epimedium Grandiflorum (Horny Goat Weed Ext.) Dry Conc 20:1"],
    "Fennel Dry Extract 5:1": ["Funiculum Vulgaris (Fennel Ext.) Dry Conc 5:1"],
    "Ferrous Bisglycinate": ["Ferrous BisGlycinate 20%"],
    "Fluticasone Propionate": ["Fluticasone Propionate Micronized"],
    "Folic Acid (Vitamin B9)": ["Folic Acid 10% Powder", "Folic Acid 10% with Maltodextrin"],
    "Garcinia Cambogia Fruit Dry Extract 6:1": ["Garcinia Quaesita (Garcinia Ext.) Dry Conc 6:1"],
    "Gentian Dry Extract 4:1": ["Gentiana Lutea (Gentian Ext.) Dry Conc 4:1"],
    "Ginger Dry Extract 5:1": ["Zingiber Officinals (Ginger Ext.) Dry Conc 5:1"],
    "Ginkgo Biloba Leaf Dry Extract 50:1": ["Gink Biloba (GinKgo Ext.) Dry Conc 50:1"],
    "Gymnema Sylvestre Leaf Dry Extract 10:1": ["Gymnema Sylvestre Leaf (Gymnema Ext.) Dry Conc 10:1"],
    "Inositol (Vitamin B8)": ["Vitamin B8 (Inositol)"],
    "Juniper Communis Fruit Dry Extract 2:1": ["Juniper"],
    "Liquorice Dry Extract 8:1": ["Glycyrrhiza Glabra (Liquorice Ext.) Dry Conc 8:1"],
    "Loperamide HCl": ["Lopromaide HCL"],
    "Magnesium Oxide Heavy": ["Magnesium Oxide (Heavy)"],
    "Methionine": ["L-Methionine"],
    "Methyl Sulphonyl Methane (MSM)": ["Methyl Sulphonyl Methan 99.4% (MSM)"],
    "Milk Thistle Dry Extract 70:1": ["Silybum Marianum (Milk Thistle Ext.) Dry Conc 70:1"],
    "Nicotinamide (Vitamin B3)": ["Vitamin B3 (Nicothinamide)"],
    "Nicotinamide Mononucleotide (NMN) Granule": ["Nicotinamide MonoNucleotide (NMN)"],
    "Omega 3 Powder (from Flax Seed Dry Extract)": ["Omega 3 (Flaxed Oil)"],
    "Ozanimod Hydrochloride": ["Ozanimod"],
    "Panax Ginseng (Asian Ginseng) Dry Extract 3:1": ["Panax Ginseng (Korean Ginseng Ext.) Dry Conc 3:1"],
    "Pyridoxine Hydrochloride (Vitamin B6)": ["Vitamin B6 (Pyridoxin Hydrochloride)"],
    "Riboflavin (Vitamin B2)": ["Vitamin B2 (Riboflavine)"],
    "Sacubitril/Valsartan 3Sodium 2.5H2O Co-crystal": ["Sacubitril Valsartan Sodium"],
    "Saw Palmetto Dry Extract 4:1": ["Serenoa Repens (Sawpalmetto Ext.) Dry Conc 4:1"],
    "Siberian Ginseng Dry Extract 5:1": ["Eleutherococuss Senticosus (Siberian Ginseng Ext.) Dry Conc 5:1"],
    "Sodium Selenite": ["Sodium Selenite Anhydrous"],
    "Thiamine Hydrochloride (Vitamin B1)": ["Vitamin B1 (Thiamin HCL)"],
    "Tribulus Terrestris Dry Extract 50:1": ["Tribulus Tresstriss Ext. Dry Conc 50:1"],
    "Ubidecarenone (Co-Q10)": ["Ubidecarenon (Co Enzyme Q10)"],
    "Vitamin D3": ["Vitamin D3 100 Dry (Cholecalciferol)", "Vitamin D3 Liquid"],
    "Apple Green Color": ["Apple Green"],
    "Brown Color": ["Brown  Color - Code 1005"],
    "Exacol Brown HT": ["Exacol Brown"],
    "Sunset Yellow": ["Sunset Yellow - FD&C Yelloe No.6"],
    "Ferric Oxide Black": ["Black Iron Oxide"],
    "Ferric Oxide Red": ["Red Iron Oxide"],
    "Ferric Oxide Yellow": ["Yellow Iron Oxide"],
    "Croscarmellose Sodium": ["Cross Carmellose Sodium"],
    "Crospovidone": ["Crosspovidone A"],
    "Lactose Monohydrate": ["Lactose Monohydrate 80", "Lactose Monohydrate 200", "Lactose monohydrate mesh 80 DC grade"],
    "Microcrystalline Cellulose PH 101": ["Micro Cryslaline Cellulose PH101 (Avicel PH101)"],
    "Microcrystalline Cellulose PH 102": ["Micro Cryslaline Cellulose PH102 (Avicel PH102)"],
    "Microcrystalline Cellulose RC591": ["Avicel RC 591"],
    "Pregelatinized Starch": ["Corn Starch 1500 Pregelatin"],
    "PVP K30": ["Poly Vinil Pyrolidone K30 (PVP K30)"],
    "PVP K90": ["Poly Vinil Pyrolidone K90 (PVP K90)"],
    "Colloidal Silicon Dioxide": ["Silica Colloidal Anhydrous (Aerosil)"],
    "Glyceryl Behenate (Compritol 888 ATO)": ["Compritol 888 ATO (Glyceryl Dibehenate)"],
    "HPMC K4M": ["Hypermellose (HPMC Benecel E4M)"],
    "Sodium Carboxymethyl Cellulose (Medium Viscosity)": ["Carboxymethyl Cellulose Sodium"],
    "Sucrose": ["Sacarose (Pharma Sugar)"],
    "Cola Flavour": ["Cola Flavour Powder"],
    "Honey Flavour": ["Honey Flovour"],
    "Lemon Flavour": ["Lemon Flovour"],
    "Mint Flavour": ["Mint Essence"],
    "Orange Flavour": ["Orange Flavour Liquid", "Orange Flavour Powder (Firmenich 05605TP0551)"],
    "Sour Cherry Flavour": ["Sour Cherry Flovour"],
    "Strawberry Flavour": ["Powder Strawberry Flovour (Firmenich)", "Strawberry Liquid Flovour",
                           "Strawberry Liquid Flovour (Firmenich 052311B)"],
    "Polyethylene Glycol 4000 (PEG 4000)": ["Poly Ethylen Glycol 4000 (PEG 4000)"],
    "Propylene Glycol": ["Propylen Glaycol"],
    "Methacrylic Acid Copolymer (Eudragit L100-55)": ["Metaacrylic Acid Copolymer (Type C) (Eudragit L100-55)"],
    "Ethanol 96%": ["Ethanol (Alcohol)"],
    "Sodium Citrate Dihydrate": ["Sodium Citrate"],
    "Benzalkonium Alcohol 50%": ["Benzalkonium Chloride"],
}

# Catalog of general tests: timed once on the 'Test catalog' sheet. code -> (name, study instrument, lab catalog names)
USP_CATALOG = collections.OrderedDict([
    ("APPEAR", ("Appearance / description and solubility", "Other / bench", ["Description", "Soluibility"])),
    ("ID_IR", ("Identification - IR spectrum", "Other / bench", [])),
    ("ID_UV", ("Identification - UV spectrum", "UV-Vis", [])),
    ("ID_CHEM", ("Identification - chemical reaction (incl. <191>)", "Other / bench", [])),
    ("TLC", ("TLC (identification or limit)", "Other / bench", ["TLC (drug)"])),
    ("LOD", ("Loss on drying", "Other / bench", ["LOD (60 min)"])),
    ("KF", ("Water - Karl Fischer", "KF titrator", ["Water content"])),
    ("ROI", ("Residue on ignition / sulfated ash", "Other / bench", ["Ash"])),
    ("LOI", ("Loss on ignition", "Other / bench", [])),
    ("ASH", ("Total / acid-insoluble ash (botanical)", "Other / bench", ["Ash"])),
    ("PH", ("pH", "Other / bench", ["pH"])),
    ("ACID", ("Acidity / alkalinity (indicator titration)", "Other / bench", ["Acidity/Alkalinity"])),
    ("OR", ("Optical rotation", "Other / bench", ["Optical rotation"])),
    ("MP", ("Melting range", "Other / bench", ["Melting point"])),
    ("SG", ("Specific gravity / density", "Other / bench", ["SPECIFIC GRAVITY"])),
    ("RI", ("Refractive index", "Other / bench", [])),
    ("VISC", ("Viscosity", "Other / bench", [])),
    ("COND", ("Conductivity", "Other / bench", ["Conductivity"])),
    ("TOC", ("Total organic carbon", "Other / bench", [])),
    ("CLAR", ("Clarity / colour of solution", "Other / bench", ["Clarity of solution", "Color of Solution"])),
    ("LIMIT", ("Limit test - wet chemistry (chloride, sulfate, iron, barium, ...)", "Other / bench", ["Chloride", "Sulfate"])),
    ("LIM_TOX", ("Arsenic / lead / mercury limit (digestion or special apparatus)", "Other / bench", [])),
    ("GRAV", ("Gravimetric limit (water-soluble, insoluble matter, residue on evaporation)", "Other / bench", ["Residue on evaporation"])),
    ("AAS", ("Element limit by AAS / ICP (Pb, Cd, Ni, Al, ...)", "AAS / ICP", [])),
    ("TITR", ("Assay or test by titration", "Other / bench", [])),
    ("UV", ("Assay or limit by UV-Vis", "UV-Vis", [])),
    ("FATS", ("Fats and oils value (acid, iodine, peroxide, saponification, hydroxyl)", "Other / bench", ["Acid value", "Peroxide value"])),
    ("PHYS", ("Powder physical test (bulk density, surface area, settling volume)", "Other / bench", ["Density Bulk", "Settling volume"])),
    ("RS_GC", ("Residual solvents by GC (<467>, <565>)", "GC", [])),
    ("PEST", ("Pesticide residues (<561>, <565>)", "GC", [])),
])

HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow = you type (same as the study)
AUTO = PatternFill("solid", fgColor="E7E6E6")    # grey = pre-filled / reference
SEC = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def norm(s):
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\(.*?\)", "", str(s).lower()))


def num(v):
    return v if isinstance(v, (int, float)) else None


def cell(v):
    return "" if v is None else str(v).replace("\n", " ").strip()


# Study wording that points at source files -> neutral wording for respondents.
NEUTRAL = [("(from database)", "(from product formulas)"), ("not in source file", "not on record"),
           ("Not in source file", "Not on record")]


def neutral(v):
    if isinstance(v, str):
        for a, b in NEUTRAL:
            v = v.replace(a, b)
    return v


def load_study():
    wb = openpyxl.load_workbook(STUDY, data_only=True)
    rows = lambda name: [tuple(neutral(v) for v in r) for r in wb[name].iter_rows(min_row=5, values_only=True) if r[0]]
    hdr = lambda name: [neutral(c.value) for c in wb[name][4]]
    return dict(tt=rows("Test_Times"), tt_h=hdr("Test_Times")[:26],
                rm=rows("RM_Lots"), rm_h=hdr("RM_Lots")[:8],
                fp=rows("Products"), fp_h=hdr("Products")[:13],
                interview=[r for r in wb["Interview"].iter_rows(min_row=5, values_only=True) if r[0]])


def load_leadtime():
    wb = openpyxl.load_workbook(LEAD, data_only=True)
    out = {}
    for ws in wb.worksheets[1:]:
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r[0]:
                continue
            g, s = num(r[12]), num(r[16])
            t = num(r[20]) if len(r) > 20 else None
            if t is None and (g is not None or s is not None):
                t = (g or 0) + (s or 0)
            out[cell(r[0])] = dict(sheet=ws.title, gt=cell(r[4]), st=cell(r[8]), g=g, s=s, t=t)
    return out


def match(study_names, lead):
    by_norm = collections.defaultdict(list)
    for n in lead:
        by_norm[norm(n)].append(n)
    res = {}
    for m in study_names:
        if m in ALIASES:
            res[m] = ("Alias - confirm", [n for n in ALIASES[m] if n in lead])
        elif by_norm.get(norm(m)):
            res[m] = ("Exact", by_norm[norm(m)])
        else:
            res[m] = ("No recorded time", [])
    return res


# rawmaterials.md names the fuzzy matcher pairs wrongly (different vitamin, likely HPMC E15 not K15M).
NOT_SAME = {"Vitamin k2", "HPMC 15"}


def load_raw():
    """Unique names from rawmaterials.md (first spelling kept) with how often each normalised name occurs."""
    seen = collections.OrderedDict()
    for line in open(RAW, encoding="utf-8"):
        if line.startswith("- ") and line[2:].strip():
            n = re.sub(r"\s+", " ", line[2:].strip())
            seen.setdefault(norm(n) or n, [n, 0])[1] += 1
    return list(seen.values())


def load_db_mapping():
    wb = openpyxl.load_workbook(DB, read_only=True, data_only=True)
    return {norm(r[0]): r[1] for r in wb["Name Mapping"].iter_rows(min_row=2, values_only=True) if r[0] and r[1]}


def match_raw(raw, study_names, lead, dbmap):
    """rawmaterials.md name -> (status, study name). Fuzzy hits are only suggestions (see ALIASES note)."""
    by_norm = {norm(m): m for m in study_names}
    via_alias = {norm(n): k for k, ns in ALIASES.items() for n in ns}
    lead_norm = {norm(n) for n in lead}
    res = []
    for name, count in raw:
        n = norm(name)
        if n in by_norm:
            st, m = "On master list", by_norm[n]
        elif n in via_alias:
            st, m = "On master list (alias)", via_alias[n]
        elif dbmap.get(n) in study_names:
            st, m = "On master list (name mapping)", dbmap[n]
        elif name not in NOT_SAME and (close := difflib.get_close_matches(n, by_norm, 1, 0.85)):
            st, m = "Suggested - confirm", by_norm[close[0]]
        elif n in lead_norm:
            st, m = "In time records only", None
        else:
            st, m = "Not on master list", None
        res.append((name, count, st, m))
    return res


def agg(names, lead, key):
    vals = [lead[n][key] for n in names if lead[n][key] is not None]
    if not vals:
        return None
    return round(max(vals) / 60, 2)   # minutes -> hours; worst supplier variant


# ---------------------------------------------------------------- xlsx helpers
def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def table(ws, r0, headers, rows, input_cols, widths=None, dv=None):
    """Write a header row at r0 and data below; input_cols are 1-based column numbers to colour yellow."""
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font, c.alignment, c.border = HEAD, Font(bold=True, color="FFFFFF"), WRAP, BOX
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border, c.alignment = BOX, WRAP
            c.fill = INPUT if j in input_cols else AUTO
    ws.row_dimensions[r0].height = 45
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    last = r0 + max(len(rows), 1) + 200   # dropdowns also cover rows added below the table
    for col, formula in (dv or {}).items():
        v = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(v)
        v.add(f"{col}{r0 + 1}:{col}{last}")


def qsheet(ws, sections, sub):
    """Question sheet: sections = [(section title, [(id, question, maps-to)])]. maps-to stays internal."""
    r = 1
    for j, (h, w) in enumerate([("#", 7), ("Question", 70), ("Answer", 60)], 1):
        c = ws.cell(r, j, h)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        ws.column_dimensions[get_column_letter(j)].width = w
    for sec, qs in sections:
        r += 1
        ws.cell(r, 1, sec).font = Font(bold=True)
        for j in range(1, 4):
            ws.cell(r, j).fill = SEC
        for qid, q, maps in qs:
            r += 1
            for j, v in enumerate([qid, q, None], 1):
                c = ws.cell(r, j, v)
                c.alignment, c.border = WRAP, BOX
                c.fill = INPUT if j == 3 else AUTO
            ws.row_dimensions[r].height = max(30, 15 * (len(q) // 80 + 1))
    ws.freeze_panes = "A2"


RESPONDENT = ("A. Respondent", [
    ("A1", "Name / ID, role, years of QC experience.", "Personnel A, B, D"),
    ("A2", "Which techniques/instruments are you qualified on (HPLC, GC, UV-Vis, AAS/ICP, KF, dissolution, titration)?", "Personnel C; Test_Times X"),
    ("A3", "Contract hours per week and annual leave weeks.", "Personnel E, F"),
    ("A4", "In a normal week, what % of your time is: testing / reviewing others' data / documentation & LIMS / "
           "meetings & training / cleaning & housekeeping / investigations? (should add to 100%)", "Personnel G, H"),
    ("A5", "Average overtime hours per month over the last 6 months, and the main reason.", "Personnel J; Evidence H"),
    ("A6", "Is any method or instrument in your area done by only one qualified person? Which?", "Report: key-person risk"),
])

TT_WIDTHS = [26, 6, 34, 26, 1, 9, 13, 10, 9, 10, 10, 10, 10, 10, 10, 9, 10, 10, 10, 10, 10, 9, 10, 16, 11, 30, 10]
TT_INPUT = set(range(6, 28)) - {16}   # F..AA editable, P (PERT) is computed
TT_DV = {"F": '"Y,N"', "G": '"HPLC,GC,UV-Vis,AAS / ICP,KF titrator,Other / bench,None"',
         "H": '"Low,Medium,High"', "W": '"Y,N"', "Y": '"High,Medium,Low"', "AA": '"Y,N"'}
OUTSOURCED = "Outsourced (Y/N)"   # after the study's A..Z, so paste-back by position is unchanged


def tt_rows(rows):
    # Keep study columns A..Z (E and P get formulas from tt_auto) and add the Outsourced column.
    return [[v if j not in (4, 15) else None for j, v in enumerate(r[:26])] + [None] for r in rows]


def tt_auto(ws, first, last):
    """Lookup key (E), PERT (P) and the Retest % format on Test times-style rows."""
    for r in range(first, last + 1):
        ws.cell(r, 5, f'=IF(A{r}="","",A{r}&"|"&D{r})')
        ws.cell(r, 16, f'=IF(N{r}="","",IF(OR(M{r}="",O{r}=""),N{r},(M{r}+4*N{r}+O{r})/6))').number_format = "0.00"
        ws.cell(r, 22).number_format = "0%"
    ws.column_dimensions["E"].hidden = True   # only used to match answers back to the study
    grey_out(ws, '$F2="N"', '$AA2="Y"')


def grey_out(ws, *conds):
    """Dark-grey the whole row when any condition (written for row 2) holds, e.g. Applicable = N."""
    rule = FormulaRule(formula=[f"OR({','.join(conds)})"], stopIfTrue=False,
                       fill=PatternFill("solid", fgColor="BFBFBF", bgColor="BFBFBF"), font=Font(color="7F7F7F"))
    ws.conditional_formatting.add(f"A2:{get_column_letter(ws.max_column)}{ws.max_row + 200}", rule)


def extra_tests_sheet(wb, s, kind, key_list=None, types=None, dv=None):
    """types: dropdown list for an editable Type column (default: Type fixed to kind)."""
    ws = wb.create_sheet("Extra tests")
    table(ws, 1, s["tt_h"] + [OUTSOURCED], [[None, kind] + [None] * 25 for _ in range(40)],
          set(range(1, 28)) - {5, 16} - (set() if types else {2}),
          TT_WIDTHS, dict(dv or TT_DV, **({"A": key_list} if key_list else {}), **({"B": types} if types else {})))
    tt_auto(ws, 2, 41)
    return ws


def overview(wb, name, area, scope_lines, sheets, stats):
    ws = wb.active
    ws.title = "Overview"
    title(ws, name, "Deliverable 2 - analyst edition. Part of the QC Laboratory Capacity, Staffing Gap and CAPEX Assessment.")
    lines = [("Area", area)] + scope_lines + [("", "")] + [("Sheet: " + a, b) for a, b in sheets if a != "Overview"] + [("", "")] \
        + [("How to fill", "Read the 'Guide' sheet first: rules, which sheets take new rows, and what every column means.")] \
        + [("Data basis", t) for t in stats] + [("", ""), ("Respondent", None), ("Date", None), ("Reviewed by (unit head)", None)]
    for i, (a, b) in enumerate(lines, 4):
        ws.cell(i, 1, a).font = Font(bold=True)
        c = ws.cell(i, 2, b)
        c.alignment = WRAP
        if a in ("Respondent", "Date", "Reviewed by (unit head)"):
            c.fill, c.border = INPUT, BOX
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110


# Column meanings for the Guide sheet, keyed by header text.
DESC = {
    "#": "Question number.", "Question": "The question.", "Answer": "Your answer - short is fine.",
    "Key (Family ID / Material)": "The item tested: material name (raw materials) or method family ID (finished products).",
    "Type": "RM = raw material, FP = finished product (Extra tests may also use PM = packaging, IPC, Water, EM, Other).",
    "Description": "Material category, or the products in the family.",
    "Test": "Test name.",
    "Lookup key (auto)": "Hidden, automatic: Key | Test joined, used to match your answers back to the study. Do not type.",
    "Applicable (Y/N)": "Y = you run this test on this item; N = you do not (do not delete the row).",
    "Instrument": "Main instrument the test occupies.",
    "Complexity": "Low / Medium / High - how demanding the test is.",
    "No. of separate methods": "Separate runs this test needs (e.g. 2 actives with separate HPLC methods = 2). Usually 1.",
    "Campaign group code": "Same code (e.g. G1) on tests run in one instrument sequence; their setup time is then counted once.",
    "Setup analyst h / sequence": "Hands-on hours to set up one sequence/run (mobile phase, standards, system suitability).",
    "Setup instrument h / sequence": "Instrument hours per sequence that are not per sample (equilibration, blanks, standards).",
    "Prep analyst h/sample – MIN": "Optional. Fastest realistic hands-on preparation per sample.",
    "Prep – MOST LIKELY": "Required. Usual hands-on preparation per sample (weighing, dissolving, diluting, filtering).",
    "Prep – MAX": "Optional. Slowest realistic preparation per sample.",
    "Prep PERT (auto)": "Automatic: (MIN + 4 x MOST LIKELY + MAX) / 6; equals MOST LIKELY if MIN or MAX is blank.",
    "Instrument run h / sample": "Unattended instrument time per sample (run time, oven/LOD time, titrator time).",
    "Dissolution bath h / sample": "Finished products only: bath time per sample.",
    "Data processing + doc h / sample": "Integration, calculation, LIMS/logbook entries per sample.",
    "Review h / sample": "Time a second person spends checking the data per sample.",
    "Samples per sequence": "Typical number of samples run together in one sequence/run.",
    "Retest / OOS %": "Share of samples tested again (type 5 for 5%).",
    "Applies to stability (Y/N)": "Finished products only: Y if the test is also run on stability samples.",
    "Respondent(s) / source": "Who gave the time, or where it comes from.",
    "Confidence": "How sure you are of the times: High / Medium / Low.",
    "Notes / difficulty": "Anything unusual, e.g. 'solution shared with Assay'.",
    OUTSOURCED: "Y = sent to an external lab. Then enter only in-house time (sampling, sending, checking the report) as Prep; leave instrument times blank.",
    "Catalog code": "The Test catalog row this test takes its times from. Change it only to point to a different catalog test.",
    "USP-NF source": "Monograph and section the test comes from.",
    "Recorded h (lab records)": "Reference: hours recorded by the lab in the past for this test set.",
    "Recorded test list": "Reference: tests listed in the lab's past records.",
    # Test catalog
    "Code": "Short code of the general test; Test times rows find their times by this code.",
    "Rows using it (auto)": "Automatic: number of Test times rows (and Monitoring rows) that use this test and are Applicable = Y.",
    "Lab catalog time (min)": "Reference: minutes in the lab's existing test catalog.",
    "Samples per sequence / run": "Typical number of samples run together.",
    "Notes": "Anything unusual.",
    # Materials
    "Material": "Material name (master list first, then names found elsewhere).",
    "Category": "Active, excipient, colorant, flavour, solvent ...",
    "# products using (from product formulas)": "Number of products whose formula uses the material.",
    "Lots / yr (now)": "Lots received and tested per year today.",
    "Lots / yr (future, blank = growth %)": "Expected lots per year in 2-3 years; blank = the study's default growth.",
    "Lots used future (auto)": "Hidden, calculated in the study. Leave empty.",
    "Supplier / pharmacopoeia": "Main supplier(s) and the pharmacopoeia tested to (USP, BP, EP, supplier spec).",
    "Suppliers (number)": "How many approved suppliers.",
    "Reduced testing? (none / ID+CoA / skip-lot)": "Whether some lots get reduced testing.",
    "Status": "Where the name came from; 'Suggested - confirm' and 'Not on master list' rows need a check.",
    "Same material as (master-list name)": "If this name is the same as a master-list material, give that name.",
    "Still tested by QC? (Y/N)": "For names not on the master list: is it still tested?",
    "Also listed as": "Other spellings of the same material.",
    "USP-NF monograph": "Monograph used to list the tests ('confirm' = matched by hand).",
    "Recorded h/lot (lab records, worst supplier)": "Reference: total recorded hours per lot.",
    "USP tests not in lab record": "Reference: tests the monograph has but the lab record does not.",
    "Lab-record tests not in USP": "Reference: tests the lab record has but the monograph does not.",
    # FP Volumes / Campaign groups
    "Product": "Product name and strength.",
    "Active(s) (from product formulas)": "Active ingredients.", "# Actives": "Number of actives.",
    "Dosage form (verify)": "Tablet, capsule, syrup ... - correct if wrong.",
    "Family ID": "Method family: strengths tested with the same methods.",
    "Release batches / yr (now)": "Release batches tested per year today.",
    "Stability samples / yr (now)": "Stability samples tested per year today.",
    "Release batches / yr (future, blank = growth %)": "Expected release batches per year in 2-3 years.",
    "Stability samples / yr (future)": "Expected stability samples per year in 2-3 years.",
    "Release used future (auto)": "Hidden, calculated in the study. Leave empty.",
    "Stability used future (auto)": "Hidden, calculated in the study. Leave empty.",
    "Stability time points / yr per batch on stability": "Pulls per year for one batch on stability.",
    "IPC samples tested by QC per batch": "In-process samples QC tests per batch.",
    "Active ingredient": "Active shared by several families.", "# families": "Number of families using it.",
    "Families": "Family IDs.", "Same column + mobile phase? (Y/N/partly)": "Can they run in one HPLC sequence?",
    "Group code": "Code to put in Test times column J (Campaign group code).",
    # Physical lab
    "IPC checks per batch (Physical lab)": "In-process checks per batch done by the Physical lab at the machine (0 = production does them all).",
    "Units tested per IPC check": "Tablets / capsules / sachets tested in one in-process check.",
    # Micro lab
    "Item (material / family)": "Raw-material name or finished-product family ID.",
    "Lots / batches tested for micro (%)": "Share of lots or batches tested for micro; blank = every one (20 = one in five).",
    "Equipment": "Main equipment the test occupies.",
    "Compendial conditions (reference)": "Reference: typical media and incubation from USP; your procedure may differ.",
    "Setup analyst h / session": "Hands-on hours per session not spent on single samples (media out, controls, cabinet cleaning).",
    "Prep + inoculation h/sample – MIN": "Optional. Fastest realistic hands-on time per sample (weighing, dilution, filtration, plating).",
    "Prep + inoculation – MOST LIKELY": "Required. Usual hands-on time per sample to prepare and inoculate.",
    "Prep + inoculation – MAX": "Optional. Slowest realistic preparation per sample.",
    "Reading + counting h / sample (all reads)": "Required. Hands-on time to read, count and subculture one sample, all reads together.",
    "Incubation days": "Days the plates / tubes stay in the incubator (not analyst time).",
    "Plates / tubes per sample": "Plates, tubes or bottles one sample puts in the incubators.",
    "BSC / LAF h / sample": "Biosafety cabinet / laminar-flow hood time per sample.",
    "Data + doc h / sample": "Calculations, logbook and LIMS entries per sample.",
    "Samples per session": "Typical number of samples handled in one session.",
    "Basis (USP-NF)": "Why the test is listed: monograph chapter, natural origin or USP <1111> for the dosage form.",
    "Programme": "Water, Environment, Personnel, Utilities, Cleaning validation ...",
    "Area / system / point group": "Area, water system or group of points with the same sampling plan.",
    "Grade / class": "Cleanroom grade or area class, if any.",
    "Points (number)": "Number of sampling points (or plates per round) in this group.",
    "Samples per point per week": "How often each point is sampled per week (daily on 5 days = 5; monthly = 0.25).",
    "Weeks per year (blank = 50)": "Weeks per year the plan runs; blank = 50.",
    "Samples / yr (auto)": "Automatic: Points x Samples per point per week x Weeks per year.",
    "Sampling round analyst h (walk + sampling)": "Hands-on hours for one sampling round, including walking and gowning.",
    "Rounds per week": "Sampling rounds per week for this group.",
    "Weekend / holiday sampling (Y/N)": "Y if samples are also taken at weekends or on holidays.",
    "Who samples (QC micro / production / engineering)": "Who takes the samples.",
    "Activity": "Support task.", "Unit": "What one unit is (batch, load, week ...).", "Units / yr": "Units per year.",
    "Analyst h / unit": "Hands-on hours per unit.", "Equipment occupied": "Equipment used (autoclave, incubator ...).",
    "Equipment h / unit": "Hours the equipment is occupied per unit.", "Done by (QC micro / other)": "Who does it.",
}

GENERAL_RULES = [
    ("Colours", "Yellow = you type. Grey = pre-filled or automatic; type over a grey value only if it is wrong. "
     "Dark-grey row = Applicable = N or Outsourced = Y (outsourced rows still take the in-house Prep time)."),
    ("Units", "Times in HOURS (0.25 = 15 min) per sample unless the header says per sequence."),
    ("Minimum effort", "Only 'Prep – MOST LIKELY' is required; MIN and MAX are optional. Blank = not known."),
    ("Hands-on vs instrument", "Analyst columns = hands-on time only. Unattended time goes in 'Instrument run'."),
    ("Shared sample solution", "When one solution serves several tests (e.g. Assay + Related substances, or pH + Clarity "
     "from one solution), enter its preparation once, on the first test. On the other tests enter only their extra "
     "hands-on time and write 'solution shared with <test>' in Notes."),
    ("Shared sequence", "Tests run in the same instrument sequence get the same Campaign group code (column J); "
     "their setup time is then counted once."),
    ("Outsourced tests", "Outsourced = Y. Enter only in-house time (sampling, sending, checking the report) as Prep; "
     "leave instrument times blank."),
    ("Columns", "Do not insert, delete or reorder columns - answers are collected by column position."),
    ("Adding rows", "Use the empty rows at the bottom of a sheet. If you need more: select the last empty row, Ctrl+C, "
     "right-click the row number below it, 'Insert Copied Cells' - colours, dropdowns and formulas come with it."),
]


def guide(ws, wb, sheets, rules, general=None):
    """sheets: [(sheet, use, may add rows)]. Then the meaning of every column of each table sheet (row 1 headers)."""
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110
    r = 1

    def head(*vals):
        nonlocal r
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        r += 1

    def line(a, b, fill=None):
        nonlocal r
        for j, v in enumerate((a, b), 1):
            c = ws.cell(r, j, v)
            c.alignment, c.border = WRAP, BOX
            if fill:
                c.fill = fill
        ws.cell(r, 1).font = Font(bold=True)
        r += 1

    head("How to fill", "")
    for a, b in (GENERAL_RULES if general is None else general) + rules:
        line(a, b)
    r += 1
    head("Sheet", "Use  ·  Can I add rows?")
    for name, use, add in sheets:
        line(name, f"{use}  ·  {add}")
    for name, *_ in sheets:
        if name in ("Overview", "Guide"):
            continue
        sh = wb[name]
        r += 1
        head(f"{name} - columns", "Meaning")
        for c in sh[1]:
            if c.value is None:
                continue
            col = c.column_letter
            line(f"{col} · {c.value}", DESC.get(c.value, ""), AUTO if sh.column_dimensions[col].hidden else None)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- RM
RM_QUESTIONS = [("A. Respondent", [RESPONDENT[1][0],
        ("A2", "Which techniques/instruments are you qualified on (HPLC, GC, UV-Vis, AAS/ICP, KF, titration, IR)?", "Personnel C; Test_Times X")]
        + RESPONDENT[1][2:5]),
    ("B. Receipt and sampling", [
        ("B1", "Who samples incoming materials (QC analyst, sampler, warehouse)? QC hours per week spent on sampling, including walking to the warehouse.", "Capacity C (non-routine)"),
        ("B2", "Which months have the most receipts, and about how many lots arrive in the peak month? (Lots per year per material go on 'Materials'.)", "Evidence B"),
        ("B3", "Which materials are re-tested at re-test date or for shelf-life extension? Lots per year?", "RM_Lots H; Capacity C"),
    ]),
    ("C. Instruments", [
        ("C1", "Which instrument do you wait for most (HPLC, GC, KF, AAS, oven)? How often per week, and how long each time?", "Evidence; Notes"),
    ]),
    ("D. Workload and service level", [
        ("D1", "Target turnaround from receipt to release (days) and actual typical turnaround. Has production ever waited for a raw material release?", "Evidence E, F, G"),
        ("D2", "How many lots are typically waiting (in quarantine, not yet tested) at month end?", "Evidence D"),
        ("D3", "Non-routine work - give number per year and hours each for: (a) new-supplier qualification (full testing of 3 lots); "
               "(b) method update and verification after a monograph change; (c) OOS investigations; (d) supplier complaints; (e) R&D or trial samples.", "Capacity C"),
    ]),
    ("E. Closing", [
        ("E1", "If you had one more analyst or one more instrument, what would change first?", "Report_Outline"),
        ("E2", "What is the biggest risk to on-time raw material release today?", "Report_Outline"),
    ]),
]

FP_QUESTIONS = [RESPONDENT,
    ("B. Batches and stability", [
        ("B1", "How many release batches does each family produce per year now, and per the production plan in 2-3 years? (enter on 'Volumes')", "Products F, I"),
        ("B2", "Stability: which conditions and time points are on the protocol (e.g. 0/3/6/9/12/18/24/36 months, 25/60 and 40/75)? How many stability samples tested per year?", "Products G, J"),
        ("B3", "Are in-process chemical tests (blend uniformity, LOD / water of granules) done by QC or by production? Samples per batch and hours? "
               "(Weight, hardness and thickness checks are on the Physical lab questionnaire.)", "Extra tests"),
        ("B4", "How many process/cleaning validation and swab samples per year? Analyst and HPLC hours each?", "Capacity C"),
    ]),
    ("C. Methods and instruments", [
        ("C1", "Which families can be run in the same HPLC sequence (same column and mobile phase)? Use the 'Campaign groups' sheet.", "Test_Times J"),
        ("C2", "Dissolution: 6 or 12 vessels? Single point or profile? HPLC or UV finish? How long is the bath occupied per sample?", "Test_Times G, R"),
        ("C3", "Content uniformity: 10 or 30 units? Is it run on the same sequence as assay?", "Test_Times U, J"),
        ("C4", "Fixed-dose combinations: does each active need its own method (e.g. ezetimibe + rosuvastatin)? Which families?", "Test_Times I"),
        ("C5", "How many days per month is each HPLC/dissolution bath down (breakdown, calibration, waiting for parts/columns)?", "Settings downtime"),
        ("C6", "Do you run HPLC overnight or at weekends unattended? What prevents it?", "Settings hours/day"),
        ("C7", "Do you wait for a free HPLC or dissolution bath? How often per week and for how long?", "Evidence; Notes"),
    ]),
    ("D. Workload and service level", [
        ("D1", "Target turnaround from sample receipt to batch release (days) and actual typical turnaround. Batches released late last year?", "Evidence E, F, G"),
        ("D2", "How many samples (release + stability) are typically waiting at month end? Stability pulls tested late?", "Evidence D, K"),
        ("D3", "Non-routine work: method validation/transfer for new products, OOS/OOT investigations, complaint samples, R&D support. Hours per year?", "Capacity C"),
        ("D4", "Which new products or line extensions arrive in the next 2-3 years?", "Products; Settings growth"),
    ]),
    ("E. Closing", [
        ("E1", "If you had one more analyst or one more HPLC, what would change first?", "Report_Outline"),
        ("E2", "What is the biggest risk to on-time batch release today?", "Report_Outline"),
    ]),
]


def load_lead_catalog():
    """Lab general-test catalog (first Leadtime sheet): name -> [minutes] (a name can occur twice)."""
    wb = openpyxl.load_workbook(LEAD, read_only=True, data_only=True)
    out = collections.defaultdict(list)
    for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if r[0] and num(r[1]) is not None:
            out[cell(r[0])].append(r[1])
    return out


def load_usp_tests():
    """data/usp_tests.csv (tools/usp_tests.py) -> material -> (monograph, mapped_by, [test rows])."""
    out = {}
    if not os.path.exists(USP_TESTS):
        return out
    for r in csv.DictReader(open(USP_TESTS, encoding="utf-8")):
        mono, how, tests = out.setdefault(r["material"], (r["monograph"], r["mapped_by"], []))
        if r["test"]:
            tests.append(r)
    return out


# Coarse test groups used to compare the USP test list with the lab's recorded test list.
GROUPS = [("ID", r"identif|\bid\b|\bir\b|infrared|\btlc\b"), ("Assay", r"assay|content of"),
          ("Impurities", r"impurit|related|purity"), ("LOD", r"loss on drying|\blod\b"),
          ("Water (KF)", r"water (determination|content)|\bkf\b|karl|moisture"),
          ("ROI / ash", r"\broi\b|residue on ignition|sul(f|ph)ated? ash|\bash\b"), ("pH", r"\bph\b"),
          ("Optical rotation", r"optical|rotation"), ("Melting", r"melting"),
          ("Heavy metals / elements", r"heavy metal|\blead\b|arsenic|mercury|cadmium|nickel|\bpb\b"),
          ("Chloride / sulfate", r"chloride|sul(f|ph)ate(?!d? ash)"), ("Clarity / colour", r"clarity|colou?r of"),
          ("Viscosity", r"viscos"), ("Specific gravity", r"specific gravity|relative density"), ("Refractive index", r"refractive")]


def groups(text):
    return {g for g, rx in GROUPS if re.search(rx, text.lower())}


def usp_groups(tests):
    g = set()
    for t in tests:
        g |= groups(t["test"])
        g |= {"ID"} if t["section"] == "Identification" else set()
        g |= {"Assay"} if t["section"] == "Assay" else set()
        g |= {"Impurities"} if t["section"] == "Impurities" and t["technique"] in ("HPLC", "GC") else set()
    return g


def cat_rules(one, own, item, differs, filt=None):
    """Guide rules for a Test catalog. own = (example test, its minutes, catalog code, rest minutes, new code);
    filt = how to find its Test times rows (default: filter on the test name)."""
    test, mins, code, rest, new = own
    return [
        ("Catalog time", f"One catalog row = the time of ONE test on ONE sample (e.g. {one}). Every Test times row "
         "with that Catalog code gets it, so each row is counted once."),
        ("One test needs its own time", f"e.g. {test} {mins} min while the rest of {code} stays {rest}: on Test catalog copy the "
         f"{code} row into a new row (Insert Copied Cells), give it a new Code (e.g. {new}) and its own times; on Test times "
         f"{filt or 'filter Test = ' + test} and change Catalog code to {new}. Those rows then follow the new row."),
        (f"One {item} differs", f"Type over that row's grey cells on Test times (e.g. {differs}). "
         "That cell then stops following the catalog."),
        ("Missing / unused tests", "Unused test: Applicable = N (do not delete the row). Missing test: add it on 'Extra tests'."),
    ]


RM_RULES = [
    ("Fill order", "1) Test catalog - each general test is timed once and copied to every material that uses it. "
     "2) Test times - filter Instrument = HPLC or GC: only those rows (and materials without a monograph) need their own times. "
     "3) Materials, then Questions."),
] + cat_rules("one chloride limit", ("Sulfate", 18, "LIMIT", 15, "SULF"), "material", "a Kjeldahl nitrogen titration")
CAT_H = ["Code", "Test", "Instrument", "Rows using it (auto)", "Lab catalog time (min)", "Setup analyst h / sequence",
         "Setup instrument h / sequence", "Prep analyst h/sample – MIN", "Prep – MOST LIKELY", "Prep – MAX",
         "Instrument run h / sample", "Data processing + doc h / sample", "Review h / sample", "Samples per sequence / run",
         OUTSOURCED, "Notes"]
CAT_LINK = {11: "F", 12: "G", 13: "H", 14: "I", 15: "J", 17: "K", 19: "L", 20: "M", 21: "N", 27: "O"}  # Test times col -> catalog col
CAT_LAST = 200   # catalog lookups cover rows 2..CAT_LAST
SPECIFIC = {"HPLC", "GC"}


def catalog_sheet(wb, rows):
    """'Test catalog' sheet; rows = [(code, test, instrument, lab catalog time)]."""
    ws = wb.create_sheet("Test catalog")
    cat_rows = [[code, name, instr, f"=COUNTIFS('Test times'!$AB:$AB,$A{i},'Test times'!$F:$F,\"Y\")", lab] + [None] * 11
                for i, (code, name, instr, lab) in enumerate(rows, 2)]
    table(ws, 1, CAT_H, cat_rows, set(range(6, 17)), [9, 44, 13, 8, 18, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 30],
          {"O": '"Y,N"'})
    grey_out(ws, '$O2="Y"')
    return ws


def link_catalog(ws, codes, links=CAT_LINK, key="AB"):
    """Test times rows with a Catalog code (column key) look their times up by code, so sorting either sheet keeps them right."""
    for i, code in enumerate(codes, 2):
        if code:
            for col, ccol in links.items():
                v = f"INDEX('Test catalog'!${ccol}$2:${ccol}${CAT_LAST},MATCH(${key}{i},'Test catalog'!$A$2:$A${CAT_LAST},0))"
                ws.cell(i, col, f'=IFERROR(IF({v}="","",{v}),"")').fill = AUTO


def build_rm(s, lead, raw, dbmap):
    mats = [r[0] for r in s["rm"]]
    m = match(mats, lead)
    rm_raw = match_raw(raw, mats, lead, dbmap)
    raw_st = collections.Counter(st for *_, st, _ in rm_raw)
    used = {n for _, names in m.values() for n in names}
    unlisted = [n for n in lead if n not in used]
    usp = load_usp_tests()
    labcat = load_lead_catalog()
    wb = Workbook()

    # ---- Test times rows (built first: the catalog needs the row counts)
    tt_study = collections.defaultdict(list)
    for r in tt_rows([r for r in s["tt"] if r[1] == "RM"]):
        tt_study[r[0]].append(r)
    tt, codes = [], collections.Counter()
    for r in s["rm"]:
        mat, cat = r[0], r[1]
        mono, how, tests = usp.get(mat, (None, "none", []))
        if tests:
            for code, test, instr, src in [("APPEAR", "Appearance / description and solubility", "Other / bench", "Lab practice (not a USP test)")] + \
                    [(t["code"], t["test"], t["technique"] if t["code"] == "" else USP_CATALOG[t["code"]][1],
                      f"{mono} · {t['section']}" + (f" <{t['chapter']}>" if t["chapter"] else "")) for t in tests]:
                t = next((x for x in tests if x["test"] == test), {})
                row = [mat, "RM", cat, test, None, t.get("applicable", "Y"), instr,
                       "Medium" if instr in SPECIFIC else "Low", 1] + [None] * 13 + ["N", None, None, t.get("note") or None]
                tt.append((row, code, [code or None, src, None, None]))
                codes[code] += bool(code)
        else:
            k, names = m[mat]
            for row in tt_study[mat]:
                test = row[3]
                key, cov = ("s", "specific tests") if test.startswith("Assay") else (("g", "general tests") if test.startswith("Other compendial") else ("t", "all tests"))
                tl = "; ".join(sorted({lead[n]["st" if key == "s" else "gt"] for n in names} - {"", "-"})) if key != "t" \
                    else "; ".join(sorted({lead[n]["gt"] + " | " + lead[n]["st"] for n in names}))
                tt.append((row, "", [None, "No USP-NF monograph - in-house / supplier spec", agg(names, lead, key), tl[:250] or None]))

    # ---- Overview + Questions
    n_mapped = sum(1 for mat in mats if usp.get(mat, (0, 0, []))[2])
    nonmaster = [x for x in rm_raw if x[2] in ("Not on master list", "In time records only")]
    stats = [f"{len(mats)} materials on the master list; {n_mapped} have a USP-NF monograph, so their tests are listed one per row "
             f"({len(tt)} material x test rows in total).",
             f"{len(USP_CATALOG)} general tests are timed once on 'Test catalog'; the material rows link to them.",
             "Materials without a monograph keep the lab's recorded time per test set as a starting point."]
    overview(wb, "Raw Material Analyst Questionnaire", "QC Chemical Laboratory - raw materials (APIs, excipients, colorants, flavours, solvents)",
             [("Who fills it", "Each analyst who tests incoming raw materials (2-3 people). Each person fills a separate copy."),
              ("Time needed", "About 2 hours: questions 20 min, test catalog 30 min, the rest on the tables.")],
             [(a, b) for a, b, _ in RM_SHEETS if a != "Overview"], stats)
    gws = wb.create_sheet("Guide")
    qsheet(wb.create_sheet("Questions"), RM_QUESTIONS, "Answer in the yellow column. Short answers are fine.")

    # ---- Materials (master list + purchase list + time records, one row per name)
    ws = wb.create_sheet("Materials")
    also = collections.defaultdict(set)
    for name, _, st, sm in rm_raw:
        if sm and st.startswith("On master list"):
            also[sm].add(name)
    rows = []
    for r in s["rm"]:
        mat = r[0]
        k, names = m[mat]
        mono, how, tests = usp.get(mat, (None, "none", []))
        rec = " ".join(lead[n]["gt"] + " " + lead[n]["st"] for n in names)
        miss = extra = None
        if tests and rec.strip(" -"):
            ug, rg = usp_groups(tests), groups(rec)
            if not any(lead[n]["st"].strip(" -") for n in names):   # no specific-test list on record
                ug -= {"Assay", "Impurities"}
            miss, extra = ", ".join(g for g, _ in GROUPS if g in ug - rg) or "-", ", ".join(g for g, _ in GROUPS if g in rg - ug) or "-"
        mono_txt = (mono + (" (confirm)" if how in ("manual", "purchase list") else "")) if tests else "None - in-house / supplier spec"
        rows.append(list(r[:5]) + [None] + list(r[6:8]) + [None, None, "Master list", None, None,
                    "; ".join(sorted((also[mat] | set(names)) - {mat}, key=str.lower))[:300] or None,
                    mono_txt, agg(names, lead, "t"), miss, extra])
    other = collections.OrderedDict()
    for n in unlisted:
        other[norm(n)] = dict(name=n, src={"time records"}, also=set(), h=round(lead[n]["t"] / 60, 2) if lead[n]["t"] else None)
    for name, _, st, sm in rm_raw:
        if st in ("Not on master list", "In time records only"):
            e = other.setdefault(norm(name), dict(name=name, src=set(), also=set(), h=None))
            e["src"].add("purchase list")
            if name != e["name"]:
                e["also"].add(name)
    sugg = [[name] + [None] * 9 + ["Suggested - confirm", sm] + [None] * 6
            for name, _, st, sm in sorted(rm_raw, key=lambda x: x[0].lower()) if st == "Suggested - confirm"]
    rest = [[e["name"]] + [None] * 9 + ["Not on master list (" + " + ".join(sorted(e["src"])) + ")", None, None,
            "; ".join(sorted(e["also"])) or None, None, e["h"], None, None]
            for e in sorted(other.values(), key=lambda e: e["name"].lower())]
    new = [[None] * 10 + ["New - add name", None, "Y"] + [None] * 5 for _ in range(10)]
    hdr = s["rm_h"] + ["Suppliers (number)", "Reduced testing? (none / ID+CoA / skip-lot)", "Status",
                       "Same material as (master-list name)", "Still tested by QC? (Y/N)", "Also listed as",
                       "USP-NF monograph", "Recorded h/lot (lab records, worst supplier)",
                       "USP tests not in lab record", "Lab-record tests not in USP"]
    table(ws, 1, hdr, rows + sugg + rest + new, {4, 5, 7, 8, 9, 10, 12, 13},
          [34, 16, 9, 10, 12, 1, 20, 26, 10, 16, 22, 30, 10, 36, 30, 11, 22, 22], {"J": '"None,ID + CoA,Skip-lot"', "M": '"Y,N"'})
    for i in range(len(rows) + len(sugg) + len(rest) + 1, len(rows) + len(sugg) + len(rest) + 11):
        ws.cell(1 + i, 1).fill = INPUT
    ws.column_dimensions["F"].hidden = True   # study formula column

    catalog_sheet(wb, [(code, name, instr, "; ".join(f"{n} {', '.join(str(v) for v in labcat[n])}" for n in lab if n in labcat) or None)
                       for code, (name, instr, lab) in USP_CATALOG.items()])

    # ---- Test times
    ws = wb.create_sheet("Test times")
    table(ws, 1, s["tt_h"] + [OUTSOURCED, "Catalog code", "USP-NF source", "Recorded h (lab records)", "Recorded test list"],
          [row + [None] + ref for row, _, ref in tt], TT_INPUT, TT_WIDTHS + [9, 40, 11, 50], TT_DV)
    tt_auto(ws, 2, len(tt) + 1)
    link_catalog(ws, [code for _, code, _ in tt])
    for col in ("R", "W"):   # finished-product columns; kept (hidden) so columns paste back by position
        ws.column_dimensions[col].hidden = True

    extra_tests_sheet(wb, s, "RM", "Materials!$A$2:$A$1000")
    guide(gws, wb, RM_SHEETS, RM_RULES)
    no_mono = [r[0] for r in s["rm"] if not usp.get(r[0], (0, 0, []))[2]]
    return wb, dict(mats=len(mats), mapped=n_mapped, rows=len(tt), cat=len(USP_CATALOG), no_mono=no_mono,
                    sugg=len(sugg), other=len(rest), raw=len(raw), unlisted=len(unlisted), m=m,
                    cats=collections.Counter(r[1] for r in s["rm"]))


RM_SHEETS = [("Overview", "Scope; your name and date.", "No."),
             ("Guide", "Rules, and what each column means.", "No."),
             ("Questions", "Answer sections A–E.", "No."),
             ("Materials", "Lots/yr now and future, suppliers, reduced testing - and whether unlisted names are still tested.",
              "Yes - 10 empty rows at the bottom for materials that are missing."),
             ("Test catalog", "Time each general test once (LOD, KF, pH, IR, titration ...).",
              "Only to give one test its own time (see 'One test needs its own time' above); do not sort or delete rows."),
             ("Test times", "One row per material x compendial test: confirm Applicable; time the HPLC/GC tests.",
              "No - set Applicable = N instead of deleting; missing tests go on 'Extra tests'. Sorting and filtering are fine."),
             ("Extra tests", "Tests you run that are missing from Test times (in-house tests, residual solvents, ICP ...).",
              "Yes - 40 empty rows; pick the material from the dropdown in column A.")]


# ---------------------------------------------------------------- FP
FP_SHEETS = [("Overview", "Scope; your name and date.", "No."),
             ("Guide", "Rules, and what each column means.", "No."),
             ("Questions", "Respondent profile, batches & stability, methods, workload.", "No."),
             ("Volumes", "Release batches and stability samples per year for each product.",
              "Yes - 10 empty rows at the bottom for new products."),
             ("Test catalog", "Time each general test once (appearance, water content, pH).",
              "Only to give one test its own time (see 'One test needs its own time' above); do not sort or delete rows."),
             ("Test times", "One row per family x test: confirm Applicable; time the HPLC tests (grey rows come from the catalog).",
              "No - set Applicable = N instead of deleting; missing tests go on 'Extra tests'. Sorting and filtering are fine."),
             ("Campaign groups", "Families sharing an active ingredient - can they share an HPLC sequence?", "Yes, at the bottom."),
             ("Extra tests", "Tests missing from the list (blend uniformity, preservative assay, delivered-dose assay ...).",
              "Yes - 40 empty rows.")]
# General finished-product tests, timed once on 'Test catalog'. code -> (name, instrument, applies to stability)
# Weight, dimensions, hardness, brittleness, friability, disintegration, fill volume and IPC checks are done by the
# Physical lab: they are on the Physical lab questionnaire (PHYS_CATALOG), not here.
FP_CATALOG = collections.OrderedDict([
    ("APPEAR", ("Appearance / description", "Other / bench", "Y")),
    ("WATER", ("Water content - KF (or LOD if that is what you run)", "KF titrator", "Y")),
    ("PH", ("pH", "Other / bench", "Y")),
])
# Dosage form -> general tests. Replaces the study's single 'Other tests (non-HPLC ...)' row per family.
FP_FORM_TESTS = {"Tablet": ["APPEAR", "WATER"],
                 "Capsule": ["APPEAR", "WATER"],
                 "Sachet": ["APPEAR", "WATER"],
                 "Nasal spray": ["APPEAR", "PH"]}

FP_RULES = [
    ("Fill order", "1) Test catalog - each general test is timed once and copied to every family that uses it. "
     "2) Test times - filter Instrument = HPLC: those rows need their own times per family. 3) Volumes, then Questions."),
] + cat_rules("one KF water determination", ("Water content of sachets", 45, "WATER", 20, "WAT_SA"), "family",
              "the water content of an effervescent product that takes much longer",
              "filter Description for the sachet families") + [
    ("Physical tests", "Average weight, diameter, thickness, hardness, brittleness, friability, disintegration, fill volume "
     "and in-process checks are done by the Physical lab - they are on the Physical lab questionnaire, not here."),
    ("Micro tests", "Microbial tests are done by the Micro lab - not part of this questionnaire."),
    ("Families", "A family = all strengths tested with the same methods. If one strength needs a different method, say so in Notes."),
    ("Campaign groups", "Families sharing an active: if they use the same column and mobile phase, give them one group "
     "code (e.g. G1) and put it in Test times column J."),
]


def campaign_candidates(fp):
    fams = collections.OrderedDict()
    for r in fp:
        if r[4]:
            f = fams.setdefault(r[4], dict(products=[], actives=set(), form=set()))
            f["products"].append(r[0])
            f["actives"].update(a.strip() for a in str(r[1] or "").split(",") if a.strip() and not a.startswith("Not on record"))
            f["form"].add(r[3] or "?")
    by_active = collections.defaultdict(list)
    for fid, f in fams.items():
        for a in f["actives"]:
            by_active[a].append(fid)
    rows = [[a, len(fids), ", ".join(fids), None, None, None]
            for a, fids in sorted(by_active.items(), key=lambda x: (-len(x[1]), x[0])) if len(fids) > 1]
    no_active = [fid for fid, f in fams.items() if not f["actives"]]
    return fams, rows, no_active


def build_fp(s):
    wb = Workbook()
    fams, camp, no_active = campaign_candidates(s["fp"])
    fp_rows, codes = [], []
    for r in tt_rows([r for r in s["tt"] if r[1] == "FP"]):
        if not r[3].startswith("Other tests"):
            fp_rows.append(r)
            codes.append(None)
            continue
        forms = fams[r[0]]["form"] if r[0] in fams else set()
        for code in dict.fromkeys(c for f in sorted(forms) for c in FP_FORM_TESTS.get(f, ["APPEAR"])):
            name, instr, stab = FP_CATALOG[code]
            row = list(r)
            row[3], row[5], row[6], row[7], row[8], row[22] = name, "Y", instr, "Low", 1, stab
            fp_rows.append(row)
            codes.append(code)
    stats = [f"{len(s['fp'])} products in {len(fams)} method families and {len(fp_rows)} family x test rows.",
             f"{len(FP_CATALOG)} general tests (appearance, water content, pH) are timed once on "
             "'Test catalog' and linked to the family rows by dosage form.",
             f"{len(camp)} active ingredients appear in more than one family - candidates for shared HPLC sequences ('Campaign groups').",
             f"{len(no_active)} families have no active ingredient in the product formulas ({', '.join(no_active)}) - name the actives in Test times Notes.",
             "No recorded finished-product test times exist yet (lab records cover raw materials only), so all times are new estimates."]
    overview(wb, "Finished Product Analyst Questionnaire", "QC Chemical Laboratory - finished products (release, stability, in-process)",
             [("Who fills it", "Each analyst who tests finished products (2-3 people). Each person fills a separate copy."),
              ("Time needed", "About 3 hours. Start with the families you test most often.")],
             [(a, b) for a, b, _ in FP_SHEETS], stats)
    gws = wb.create_sheet("Guide")
    qsheet(wb.create_sheet("Questions"), FP_QUESTIONS, "Answer in the yellow column. Short answers are fine.")

    ws = wb.create_sheet("Volumes")
    rows = [list(r[:10]) + [None, None, r[12], None, None] for r in s["fp"]] + [[None] * 15 for _ in range(10)]
    table(ws, 1, s["fp_h"] + ["Stability time points / yr per batch on stability", "IPC samples tested by QC per batch"],
          rows, {4, 5, 6, 7, 8, 9, 10, 13, 14, 15}, [30, 30, 8, 12, 8, 10, 10, 12, 12, 12, 1, 1, 30, 14, 14],
          {"H": '"Marketed,Pipeline / no formula,Discontinued"'})
    for i in range(len(s["fp"]) + 2, len(rows) + 2):
        for j in (1, 2, 3):
            ws.cell(i, j).fill = INPUT
    for col in ("K", "L"):   # study formula columns
        ws.column_dimensions[col].hidden = True

    ws = catalog_sheet(wb, [(code, name, instr, None) for code, (name, instr, _) in FP_CATALOG.items()])
    ws.column_dimensions["E"].hidden = True   # no lab catalog times for finished products

    ws = wb.create_sheet("Test times")
    table(ws, 1, s["tt_h"] + [OUTSOURCED, "Catalog code"], [row + [c] for row, c in zip(fp_rows, codes)],
          TT_INPUT, TT_WIDTHS + [9], TT_DV)
    tt_auto(ws, 2, len(fp_rows) + 1)
    link_catalog(ws, codes)

    ws = wb.create_sheet("Campaign groups")
    table(ws, 1, ["Active ingredient", "# families", "Families", "Same column + mobile phase? (Y/N/partly)",
                  "Group code", "Notes"], camp, {4, 5, 6}, [36, 10, 40, 16, 10, 40], {"D": '"Y,N,Partly"'})

    extra_tests_sheet(wb, s, "FP")
    guide(gws, wb, FP_SHEETS, FP_RULES)
    return wb, dict(products=len(s["fp"]), fams=len(fams), no_active=no_active, rows=len(fp_rows), camp=camp,
                    forms=collections.Counter(r[3] or "not stated" for r in s["fp"]))


# ---------------------------------------------------------------- Physical lab
# Physical tests on finished products, timed once on 'Test catalog'. code -> (name, instrument, applies to stability)
PHYS_CATALOG = collections.OrderedDict([
    ("AVG_WT", ("Average weight / uniformity of mass", "Balance", "N")),
    ("DIAM", ("Diameter (or length / width)", "Tablet tester", "N")),
    ("THICK", ("Thickness", "Tablet tester", "N")),
    ("HARD", ("Hardness / breaking force", "Tablet tester", "Y")),
    ("BRITTLE", ("Brittleness", "Other / bench", "N")),
    ("FRIAB", ("Friability", "Friabilator", "N")),
    ("DISINT", ("Disintegration", "Disintegration tester", "Y")),
    ("IPC", ("In-process check at the machine (one check, all its tests)", "Tablet tester", "N")),
    ("FILL", ("Fill volume / deliverable volume", "Balance", "N")),
    ("SPRAY", ("Spray weight per actuation (shot weight)", "Balance", "Y")),
])
PHYS_FORM_TESTS = {"Tablet": ["AVG_WT", "DIAM", "THICK", "HARD", "BRITTLE", "FRIAB", "DISINT", "IPC"],
                   "Capsule": ["AVG_WT", "DISINT", "IPC"],
                   "Sachet": ["AVG_WT", "IPC"],
                   "Nasal spray": ["FILL", "SPRAY"]}
PHYS_INSTR = '"Balance,Tablet tester,Caliper / micrometer,Friabilator,Disintegration tester,Dissolution bath,' \
             'Moisture analyser,Sieve shaker,Density tester,Leak tester,Other / bench,None"'
PHYS_DV = dict(TT_DV, G=PHYS_INSTR)

PHYS_SHEETS = [("Overview", "Scope; your name and date.", "No."),
               ("Guide", "Rules, and what each column means.", "No."),
               ("Questions", "Respondent profile, samples & IPC, methods, workload.", "No."),
               ("Volumes", "In-process checks per batch for each product (and release / stability numbers if you know them).",
                "Yes - 10 empty rows at the bottom for new products."),
               ("Test catalog", "Time each physical test once (weight, diameter, thickness, hardness, brittleness, friability, disintegration ...).",
                "Only to give one test its own time (see 'One test needs its own time' above); do not sort or delete rows."),
               ("Test times", "One row per family x physical test: confirm Applicable (grey rows come from the catalog).",
                "No - set Applicable = N instead of deleting; missing tests go on 'Extra tests'. Sorting and filtering are fine."),
               ("Extra tests", "Tests missing from the list: raw-material (particle size, bulk / tapped density), packaging "
                "(dimensions, leak test), dissolution bath work ...", "Yes - 40 empty rows; choose the Type (FP, RM, PM, IPC).")]

PHYS_QUESTIONS = [("A. Respondent", [RESPONDENT[1][0],
        ("A2", "Which instruments are you qualified on (balance, tablet tester, friabilator, disintegration tester, dissolution bath, "
               "moisture analyser, sieve shaker, density tester, leak tester)?", "Personnel C; Test_Times X")]
        + RESPONDENT[1][2:]),
    ("B. Samples and in-process checks", [
        ("B1", "Who does the in-process checks during compression / encapsulation / filling: the Physical lab, production operators, or both? "
               "How many checks per batch and how many units per check? (enter on 'Volumes')", "Volumes"),
        ("B2", "Do you walk to the production line for in-process checks, or are samples brought to the lab? Minutes per round trip?", "Test catalog IPC"),
        ("B3", "Stability: which physical tests are run at stability pulls (hardness, disintegration, ...)? How many stability samples per year?", "Test_Times W"),
        ("B4", "Who checks appearance / description of finished products and raw materials: the Physical or the Chemical lab?", "Report: test split"),
        ("B5", "Does the Physical lab test raw materials (particle size, bulk / tapped density, sieve analysis, flow) or packaging materials "
               "(dimensions, foil thickness, leak test)? Which tests, and how many samples per year? (add them on 'Extra tests')", "Extra tests"),
    ]),
    ("C. Methods and instruments", [
        ("C1", "Are weight, thickness, diameter and hardness measured on one automatic tablet tester in one run, or one by one "
               "(balance, caliper, hardness tester)?", "Test catalog"),
        ("C2", "Brittleness: how is it measured, on which instrument, and how is it different from hardness and friability?", "Test catalog BRITTLE"),
        ("C3", "Disintegration: how many baskets / stations? How long for uncoated, film-coated and enteric-coated tablets? How many runs at once?", "Test catalog DISINT"),
        ("C4", "Dissolution: does the Physical lab run the bath and take the samples (with UV / HPLC finish in the Chemical lab)? If yes, add it on 'Extra tests'.", "Extra tests"),
        ("C5", "How many days per month is each instrument down (breakdown, calibration, waiting for repair)?", "Settings downtime"),
        ("C6", "Which instrument do you wait for most? How often per week and for how long?", "Evidence; Notes"),
    ]),
    ("D. Workload and service level", [
        ("D1", "Target time for an in-process result (minutes) and for release results (days); actual typical times. Has production waited for you?", "Evidence E, F, G"),
        ("D2", "How many samples are typically waiting at month end? Which months are busiest (production campaigns)?", "Evidence B, D"),
        ("D3", "Non-routine work - number per year and hours each: (a) instrument calibration / mechanical checks (disintegration tester, friabilator, "
               "hardness tester, dissolution bath, daily balance checks); (b) process validation batches with extra sampling; (c) OOS investigations; "
               "(d) complaint samples; (e) R&D or trial batches.", "Capacity C"),
        ("D4", "Which new products in the next 2-3 years need new physical tests (chewable, orodispersible, effervescent, coated, modified release)?", "Products; Settings growth"),
    ]),
    ("E. Closing", [
        ("E1", "If you had one more analyst or one more instrument, what would change first?", "Report_Outline"),
        ("E2", "What is the biggest risk to on-time in-process and release results today?", "Report_Outline"),
    ]),
]

PHYS_GENERAL = [r for r in GENERAL_RULES if r[0] not in ("Hands-on vs instrument", "Shared sample solution", "Shared sequence")] + [
    ("Hands-on vs instrument", "Analyst columns = hands-on time only. Unattended time (disintegration run, friabilator drum, "
     "dissolution bath) goes in 'Instrument run'.")]
PHYS_RULES = [
    ("Fill order", "1) Test catalog - each physical test is timed once and copied to every family that uses it. "
     "2) Test times - confirm Applicable per family; type over a grey value only where one family differs. "
     "3) Volumes (in-process checks), then Questions."),
] + cat_rules("one friability test", ("Disintegration of film-coated tablets", 40, "DISINT", 20, "DIS_FC"), "family",
              "disintegration of an enteric-coated tablet that takes much longer",
              "filter Test = Disintegration, pick the film-coated families") + [
    ("Measured together", "If one tablet tester measures weight, thickness, diameter and hardness in one run, put the whole run on HARD "
     "and enter 0 on DIAM and THICK (and AVG_WT if it is part of the run), with 'measured with HARD' in Notes. Each minute is then counted once."),
    ("In-process checks", "The IPC row is ONE check at the machine: all tests in it, plus walking to the line. "
     "Checks per batch go on 'Volumes'."),
    ("Other labs", "Assay, dissolution finish (HPLC / UV), water content and pH are on the Chemical lab questionnaires; "
     "microbial tests on the Micro lab questionnaire. If the Physical lab runs the dissolution bath, add it on 'Extra tests'."),
    ("Volumes", "Release batches and stability samples are also asked on the FP questionnaire: fill them only if you know them. "
     "The in-process columns are what this sheet needs most."),
    ("Families", "A family = all strengths tested with the same methods. If one strength needs a different method, say so in Notes."),
]


def build_phys(s):
    wb = Workbook()
    fams, _, _ = campaign_candidates(s["fp"])
    rows, codes = [], []
    for r in tt_rows([r for r in s["tt"] if r[1] == "FP" and r[3].startswith("Other tests")]):
        forms = fams[r[0]]["form"] if r[0] in fams else set()
        for code in dict.fromkeys(c for f in sorted(forms) for c in PHYS_FORM_TESTS.get(f, ["AVG_WT"])):
            name, instr, stab = PHYS_CATALOG[code]
            row = list(r)
            row[3], row[5], row[6], row[7], row[8], row[22] = name, "Y", instr, "Low", 1, stab
            rows.append(row)
            codes.append(code)
    forms = collections.Counter(r[3] or "not stated" for r in s["fp"])
    stats = [f"{len(s['fp'])} products in {len(fams)} method families; dosage forms: "
             + ", ".join(f"{c} {n}" for n, c in forms.most_common()) + ".",
             f"{len(rows)} family x test rows; the {len(PHYS_CATALOG)} physical tests are timed once on 'Test catalog' and linked by dosage form.",
             "No recorded physical test times exist yet, so all times are new estimates."]
    overview(wb, "Physical Lab Analyst Questionnaire", "QC Physical Laboratory - finished products (release, stability, in-process checks), "
             "plus any raw-material and packaging physical tests",
             [("Who fills it", "Each analyst who runs physical tests (2-3 people). Each person fills a separate copy."),
              ("Time needed", "About 1.5 hours: questions 20 min, test catalog 30 min, the rest on the tables.")],
             [(a, b) for a, b, _ in PHYS_SHEETS], stats)
    gws = wb.create_sheet("Guide")
    qsheet(wb.create_sheet("Questions"), PHYS_QUESTIONS, "Answer in the yellow column. Short answers are fine.")

    ws = wb.create_sheet("Volumes")
    vol = [list(r[:10]) + [None, None, r[12], None, None] for r in s["fp"]] + [[None] * 15 for _ in range(10)]
    table(ws, 1, s["fp_h"] + ["IPC checks per batch (Physical lab)", "Units tested per IPC check"],
          vol, {4, 5, 6, 7, 8, 9, 10, 13, 14, 15}, [30, 30, 8, 12, 8, 10, 10, 12, 12, 12, 1, 1, 30, 14, 14],
          {"H": '"Marketed,Pipeline / no formula,Discontinued"'})
    for i in range(len(s["fp"]) + 2, len(vol) + 2):
        for j in (1, 2, 3):
            ws.cell(i, j).fill = INPUT
    for col in ("K", "L"):   # study formula columns
        ws.column_dimensions[col].hidden = True

    ws = catalog_sheet(wb, [(code, name, instr, None) for code, (name, instr, _) in PHYS_CATALOG.items()])
    ws.column_dimensions["E"].hidden = True   # no lab catalog times for physical tests
    ws.add_data_validation(dv := DataValidation(type="list", formula1=PHYS_INSTR, allow_blank=True))
    dv.add(f"C2:C{CAT_LAST}")

    ws = wb.create_sheet("Test times")
    table(ws, 1, s["tt_h"] + [OUTSOURCED, "Catalog code"], [row + [c] for row, c in zip(rows, codes)],
          TT_INPUT, TT_WIDTHS + [9], PHYS_DV)
    tt_auto(ws, 2, len(rows) + 1)
    link_catalog(ws, codes)
    for col in ("J", "R"):   # HPLC campaign group and dissolution bath: not used here, kept for paste-back by position
        ws.column_dimensions[col].hidden = True

    extra_tests_sheet(wb, s, "FP", types='"FP,RM,PM,IPC"', dv=PHYS_DV)
    guide(gws, wb, PHYS_SHEETS, PHYS_RULES, PHYS_GENERAL)
    return wb, dict(products=len(s["fp"]), fams=len(fams), rows=len(rows), forms=forms)


# ---------------------------------------------------------------- Micro lab
# Micro tests and monitoring samples, timed once on 'Test catalog'. code -> (name, equipment, compendial conditions)
MIC_CATALOG = collections.OrderedDict([
    ("ENUM", ("Microbial enumeration - TAMC + TYMC (<61> / <2021>)", "BSC + incubators", "TAMC 30-35 °C, 3-5 d; TYMC 20-25 °C, 5-7 d")),
    ("ECOLI", ("Absence of Escherichia coli (<62> / <2022>)", "BSC + incubators", "Enrichment 30-35 °C 18-24 h; MacConkey broth 42-44 °C 24-48 h; MacConkey agar 30-35 °C 18-72 h")),
    ("SALM", ("Absence of Salmonella (<62> / <2022>)", "BSC + incubators", "Enrichment 30-35 °C 18-24 h; RVS broth 30-35 °C 18-24 h; XLD agar 30-35 °C 18-48 h")),
    ("PSAE", ("Absence of Pseudomonas aeruginosa (<62>)", "BSC + incubators", "Enrichment 30-35 °C 18-24 h; cetrimide agar 30-35 °C 18-72 h")),
    ("SAUR", ("Absence of Staphylococcus aureus (<62>)", "BSC + incubators", "Enrichment 30-35 °C 18-24 h; mannitol salt agar 30-35 °C 18-72 h")),
    ("BTGN", ("Bile-tolerant Gram-negative bacteria (<62> / <2021>)", "BSC + incubators", "Pre-incubation 20-25 °C ~2 h; EE broth 30-35 °C 24-48 h; VRBG agar 30-35 °C 18-24 h")),
    ("CALB", ("Absence of Candida albicans (<62>)", "BSC + incubators", "Sabouraud broth 30-35 °C 3-5 d; Sabouraud agar 30-35 °C 24-48 h")),
    ("PET", ("Antimicrobial effectiveness / preservative efficacy (<51>)", "BSC + incubators", "5 test organisms; 20-25 °C; reads at 14 and 28 d")),
    ("BET", ("Bacterial endotoxins (<85>)", "Endotoxin reader / heating block", "Gel-clot 37 °C 60 min, or kinetic method")),
    ("STER", ("Sterility (<71>)", "Isolator / cleanroom + incubators", "14 d incubation (FTM 30-35 °C, SCDM 20-25 °C)")),
    ("WATER", ("Water - total count by membrane filtration (<1231>)", "Filtration manifold + incubators", "R2A 20-25 or 30-35 °C 5-7 d, or PCA 30-35 °C 3-5 d")),
    ("WAT_SPEC", ("Water - specified organisms / coliforms (if in your specification)", "Filtration manifold + incubators", "Per your specification")),
    ("EM_SET", ("Environmental monitoring - settle plate", "Incubators", "Per your EM procedure (e.g. 20-25 °C then 30-35 °C, 5-7 d)")),
    ("EM_CON", ("Environmental monitoring - contact plate / surface swab", "Incubators", "Per your EM procedure")),
    ("EM_AIR", ("Environmental monitoring - active air sample", "Air sampler + incubators", "Per your EM procedure")),
    ("PERS", ("Personnel monitoring - glove / gown plates", "Incubators", "Per your EM procedure")),
    ("GAS", ("Compressed air / gas - microbial sample", "Gas sampler + incubators", "Per your procedure")),
    ("SWAB", ("Cleaning validation - microbial swab / rinse", "Incubators", "Per your procedure")),
])
MIC_TT_CODES = ("ENUM", "ECOLI", "SALM", "PSAE", "SAUR", "BTGN", "CALB", "PET", "BET", "STER")
MIC_CAT_H = ["Code", "Test", "Equipment", "Compendial conditions (reference)", "Rows using it (auto)", "Setup analyst h / session",
             "Prep + inoculation h/sample – MIN", "Prep + inoculation – MOST LIKELY", "Prep + inoculation – MAX",
             "Reading + counting h / sample (all reads)", "Incubation days", "Plates / tubes per sample", "BSC / LAF h / sample",
             "Data + doc h / sample", "Review h / sample", "Samples per session", OUTSOURCED, "Notes"]
MIC_TT_H = ["Item (material / family)", "Type", "Description", "Test", "Applicable (Y/N)", "Lots / batches tested for micro (%)",
            "Setup analyst h / session", "Prep + inoculation h/sample – MIN", "Prep + inoculation – MOST LIKELY",
            "Prep + inoculation – MAX", "Prep PERT (auto)", "Reading + counting h / sample (all reads)", "Incubation days",
            "Plates / tubes per sample", "BSC / LAF h / sample", "Data + doc h / sample", "Review h / sample",
            "Samples per session", "Retest / OOS %", "Applies to stability (Y/N)", OUTSOURCED, "Respondent(s) / source",
            "Confidence", "Notes / difficulty", "Catalog code", "Basis (USP-NF)"]
MIC_TT_WIDTHS = [30, 6, 26, 34, 9, 10, 10, 10, 10, 10, 9, 10, 9, 9, 9, 10, 9, 9, 9, 10, 10, 16, 11, 30, 9, 40]
MIC_LINK = {7: "F", 8: "G", 9: "H", 10: "I", 12: "J", 13: "K", 14: "L", 15: "M", 16: "N", 17: "O", 18: "P", 21: "Q"}
MIC_DV = {"E": '"Y,N"', "T": '"Y,N"', "U": '"Y,N"', "W": '"High,Medium,Low"'}
# Natural-origin materials without a USP-NF micro test: pre-set to Y (botanicals: USP <2023>).
NATURAL = r"extract|ginseng|beet root|premix|gelatin|omega 3 powder"

MIC_SHEETS = [("Overview", "Scope; your name and date.", "No."),
              ("Guide", "Rules, and what each column means.", "No."),
              ("Questions", "Respondent profile, samples & programmes, media and equipment, workload.", "No."),
              ("Test catalog", "Time each micro test once (TAMC / TYMC, each specified organism, water, EM plates ...).",
               "Only to give one test its own time (see 'One test needs its own time' above); do not sort or delete rows."),
              ("Test times", "One row per raw material / product family x micro test: confirm Applicable and the % of lots tested.",
               "No - set Applicable = N instead of deleting; missing tests go on 'Extra tests'. Sorting and filtering are fine."),
              ("Monitoring", "Water, environmental, personnel, gas and swab programmes: points and samples per week.",
               "Yes - 15 empty rows for other areas or systems."),
              ("Support work", "Media, growth promotion, cultures, autoclave, identification, investigations ...: units per year and hours.",
               "Yes - 10 empty rows."),
              ("Extra tests", "Tests missing from Test times (other organisms, in-house tests, stability-only tests ...).",
               "Yes - 40 empty rows.")]

MIC_QUESTIONS = [("A. Respondent", [RESPONDENT[1][0],
        ("A2", "Which work are you qualified for (plate count, membrane filtration, specified organisms, preservative efficacy, water and "
               "EM sampling, media preparation, autoclave, identification, endotoxins, sterility)?", "Personnel C"),
        ] + RESPONDENT[1][2:5] + [
        ("A6", "Is any test or task done by only one qualified person? Is there a weekend / holiday rota for water or EM sampling, and how many people does it need?", "Report: key-person risk, coverage"),
    ]),
    ("B. Samples and programmes", [
        ("B1", "Raw materials: is every lot tested for micro, or a reduced frequency (e.g. first 3 lots, then 1 in 5), or CoA only? "
               "Which materials? (column F on 'Test times')", "Test times F"),
        ("B2", "Finished products: every batch or reduced? At which stability time points are micro tests done?", "Test times F, T"),
        ("B3", "Water: how many purified-water points, how often is each sampled, and who takes the samples? (enter on 'Monitoring')", "Monitoring"),
        ("B4", "Environmental and personnel monitoring: which rooms, which plate types, how often, and are weekends included? (enter on 'Monitoring')", "Monitoring"),
        ("B5", "Are sterility, endotoxin or other aseptic tests done at all (for products, water or suppliers)? If yes: samples per year and equipment.", "Extra tests"),
    ]),
    ("C. Media and equipment", [
        ("C1", "Media: prepared in-house or bought ready-made? Batches per week; growth promotion on every lot?", "Support work"),
        ("C2", "Incubators: how many, at which temperatures, how many plates each? Are they ever full? In which months?", "Report: incubator capacity"),
        ("C3", "Autoclave(s): how many, cycles per day, shared with waste decontamination? Breakdowns in the last year?", "Report: autoclave"),
        ("C4", "Biosafety cabinets / laminar-flow hoods: how many? Do you wait for a free cabinet? How often and how long?", "Evidence; Notes"),
        ("C5", "Identification of isolates: done in-house (Gram stain, biochemical kits, instrument) or sent out? How many per year?", "Support work"),
        ("C6", "How many days per month is each key item (incubator, autoclave, BSC, air sampler) down?", "Settings downtime"),
    ]),
    ("D. Workload and service level", [
        ("D1", "Turnaround from sample receipt to micro result (days) for raw materials and products. Does the micro result set the release date? "
               "Batches released late because of micro last year?", "Evidence E, F, G"),
        ("D2", "How many samples are typically waiting or in incubation at month end? Which months are busiest?", "Evidence B, D"),
        ("D3", "Non-routine work - number per year and hours each: (a) method suitability for new materials / products; (b) OOS and EM excursion "
               "investigations; (c) disinfectant efficacy studies; (d) incubator mapping, autoclave and BSC qualification; (e) audits. "
               "(enter on 'Support work')", "Support work"),
        ("D4", "Which new products or materials in the next 2-3 years need micro testing (liquids, probiotics, botanicals, preserved products)?", "Products; Settings growth"),
    ]),
    ("E. Closing", [
        ("E1", "If you had one more microbiologist, or one more incubator / autoclave / cabinet, what would change first?", "Report_Outline"),
        ("E2", "What is the biggest risk to on-time micro results today?", "Report_Outline"),
    ]),
]

MIC_GENERAL = [r for r in GENERAL_RULES if r[0] in ("Colours", "Outsourced tests", "Columns", "Adding rows")] + [
    ("Units", "Times in HOURS (0.25 = 15 min) per sample unless the header says per session. Incubation in DAYS."),
    ("Minimum effort", "Only 'Prep + inoculation – MOST LIKELY' and 'Reading + counting' are required; MIN and MAX are optional. Blank = not known."),
    ("Incubation is not analyst time", "Analyst columns = hands-on time only (preparing, inoculating, reading, counting). Incubation goes in "
     "'Incubation days': it fills incubators, not people. 'Reading + counting' covers all reads of one sample (e.g. day 3 and day 5)."),
    ("Session", "A session = one set-up at the cabinet (media out, controls, cleaning before and after). Its setup time is counted once "
     "for all samples of that session."),
]
MIC_RULES = [
    ("Fill order", "1) Test catalog - each micro test is timed once and copied to every row that uses it. "
     "2) Test times - confirm Applicable and the % of lots / batches tested. 3) Monitoring and Support work. 4) Questions."),
] + cat_rules("one TAMC + TYMC test", ("TAMC + TYMC of botanical extracts", 60, "ENUM", 30, "ENUM_BOT"), "material or family",
              "a product that needs extra dilution or neutralisation", "filter Test = Microbial enumeration, pick the extract rows") + [
    ("Volumes", "Lots and batches per year are collected on the RM and FP questionnaires. Here give only the share tested for micro "
     "in column F (blank = every lot / batch; 20 = one lot in five)."),
    ("Pre-filled tests", "Raw materials: the micro tests in their USP-NF monograph; natural-origin materials without one (extracts, "
     "premixes, gelatin) are pre-set to Y, the rest to N. Products: by dosage form (USP <1111>), plus Salmonella where the product "
     "contains a botanical. Correct anything that differs from your specifications."),
    ("Sterile work", "No sterile product is on the product list, so sterility and endotoxin rows are pre-set to N. If you run them, set Y and answer B5."),
    ("Families", "A family = all strengths tested with the same methods. If one strength needs a different method, say so in Notes."),
]

MON_H = ["Programme", "Area / system / point group", "Grade / class", "Catalog code", "Points (number)", "Samples per point per week",
         "Weeks per year (blank = 50)", "Samples / yr (auto)", "Sampling round analyst h (walk + sampling)", "Rounds per week",
         "Weekend / holiday sampling (Y/N)", "Who samples (QC micro / production / engineering)", "Notes"]
MON_ROWS = [("Water", "Purified water - generation / storage tank", "WATER"),
            ("Water", "Purified water - distribution loop, points of use", "WATER"),
            ("Water", "Purified water - specified organisms (if in your specification)", "WAT_SPEC"),
            ("Water", "Potable / feed water", "WATER"),
            ("Environment", "Production rooms (dispensing, granulation, compression, coating, filling, packaging)", "EM_SET"),
            ("Environment", "Production rooms (dispensing, granulation, compression, coating, filling, packaging)", "EM_CON"),
            ("Environment", "Production rooms (dispensing, granulation, compression, coating, filling, packaging)", "EM_AIR"),
            ("Environment", "QC micro lab (cabinets, testing room)", "EM_SET"),
            ("Environment", "QC micro lab (cabinets, testing room)", "EM_CON"),
            ("Personnel", "Production operators / micro analysts", "PERS"),
            ("Utilities", "Compressed air / nitrogen in product contact", "GAS"),
            ("Cleaning validation", "Equipment swabs / rinses (microbial)", "SWAB")]
SUP_H = ["Activity", "Unit", "Units / yr", "Analyst h / unit", "Equipment occupied", "Equipment h / unit", "Done by (QC micro / other)", "Notes"]
SUP_ROWS = [("Culture media preparation (in-house)", "batch"), ("Ready-made media - receipt and check", "lot"),
            ("Growth promotion test", "media lot"), ("Reference cultures - maintenance and subculture", "week"),
            ("Autoclave - media sterilisation", "load"), ("Autoclave - waste decontamination", "load"),
            ("Incubator / fridge temperature checks and cleaning", "week"), ("Lab cleaning and disinfection (incl. cabinets)", "week"),
            ("Positive / negative controls (if not in the test times)", "session"), ("Method suitability - new material or product", "item"),
            ("Identification of isolates (Gram stain, biochemical kit / instrument)", "isolate"),
            ("Investigations - OOS, out-of-limit, EM excursions", "event"), ("Disinfectant efficacy study", "study"),
            ("Equipment qualification (incubator mapping, autoclave, cabinet certification)", "item"),
            ("Gowning / aseptic technique qualification", "person"), ("Training given and received", "week")]


def load_usp_micro():
    """data/usp_micro.csv (tools/usp_tests.py) -> material -> [(code, chapter, applicable, note)]."""
    out = collections.defaultdict(list)
    if os.path.exists(USP_MICRO):
        for r in csv.DictReader(open(USP_MICRO, encoding="utf-8")):
            out[r["material"]].append((r["code"], r["chapter"], r["applicable"], r["note"]))
    return out


def build_micro(s):
    wb = Workbook()
    usp = load_usp_micro()
    fams, _, _ = campaign_candidates(s["fp"])
    desc = {r[0]: r[2] for r in s["tt"] if r[1] == "FP"}
    rows, codes = [], []

    def add(item, kind, description, code, appl, stab, note, basis):
        rows.append([item, kind, description, MIC_CATALOG[code][0], appl] + [None] * 14 + [stab, None, None, None, note, code, basis])
        codes.append(code)

    n_usp = n_nat = 0
    for r in s["rm"]:
        mat, cat = r[0], r[1]
        if usp.get(mat):
            n_usp += 1
            for code, ch, appl, note in usp[mat]:
                if code in ("BET", "STER"):   # no sterile product on the product list
                    appl, note = "N", note or "No sterile product on the product list - set Y if you run it"
                add(mat, "RM", cat, code, appl, None, note or None, "USP-NF monograph" + (f" <{ch.replace(',', '>, <')}>" if ch else ""))
        elif re.search(NATURAL, mat.lower()):
            n_nat += 1
            for code in ("ENUM", "ECOLI", "SALM"):
                add(mat, "RM", cat, code, "Y", None, None, "Natural origin, no USP-NF micro test - confirm (USP <2023>)")
        else:
            add(mat, "RM", cat, "ENUM", "N", None, None, "No micro test in USP-NF - set Y if your specification has one")
    n_rm = len(rows)
    for fid, f in fams.items():
        botanical, seen = any(re.search(NATURAL, a.lower()) for a in f["actives"]), set()
        for form in sorted(f["form"]):
            if form == "Nasal spray":
                tests = [("ENUM", "Y"), ("SAUR", "N"), ("PSAE", "N"), ("PET", "Y")]
                basis = "USP <1111> nasal use: TAMC, TYMC, absence of S. aureus and P. aeruginosa; <51> for a preserved product"
            else:
                tests = [("ENUM", "Y"), ("ECOLI", "N")] + ([("SALM", "N")] if botanical else [])
                basis = "USP <1111> oral solid: TAMC, TYMC, absence of E. coli" + ("; <2023> botanical: absence of Salmonella" if botanical else "")
            for code, stab in tests:
                if code not in seen:
                    seen.add(code)
                    add(fid, "FP", desc.get(fid, form), code, "Y", stab,
                        "Usually stability / validation batches only - give the % in column F" if code == "PET" else None, basis)

    by_code = collections.Counter(c for c, r in zip(codes, rows) if r[1] == "RM" and r[4] == "Y")
    stats = [f"{len(s['rm'])} raw materials on the master list: {n_usp} have a micro test in their USP-NF monograph, {n_nat} more are of "
             "natural origin (extracts, premixes, gelatin) and pre-set to Y; the rest have one enumeration row set to N.",
             f"{len(fams)} finished-product families ({len(s['fp'])} products) with tests by dosage form.",
             f"{len(rows)} item x test rows; {len(MIC_CATALOG)} micro tests and monitoring samples are timed once on 'Test catalog'.",
             "Water, environmental and personnel monitoring go on 'Monitoring'; media, cultures, autoclave and other support work on 'Support work'.",
             "No recorded micro test times exist yet, so all times are new estimates."]
    overview(wb, "Micro Lab Analyst Questionnaire", "QC Microbiology Laboratory - raw materials, finished products, water and environmental monitoring",
             [("Who fills it", "Each microbiologist (2-3 people). Each person fills a separate copy."),
              ("Time needed", "About 2 hours: questions 20 min, test catalog 30 min, the rest on the tables.")],
             [(a, b) for a, b, _ in MIC_SHEETS], stats)
    gws = wb.create_sheet("Guide")
    qsheet(wb.create_sheet("Questions"), MIC_QUESTIONS, "Answer in the yellow column. Short answers are fine.")

    ws = wb.create_sheet("Test catalog")
    cat = [[code, name, eq, cond, f"=COUNTIFS('Test times'!$Y:$Y,$A{i},'Test times'!$E:$E,\"Y\")+COUNTIF(Monitoring!$D:$D,$A{i})"]
           + [None] * 13 for i, (code, (name, eq, cond)) in enumerate(MIC_CATALOG.items(), 2)]
    table(ws, 1, MIC_CAT_H, cat, set(range(6, 19)), [10, 44, 22, 40, 8, 10, 10, 10, 10, 10, 9, 9, 9, 10, 9, 9, 10, 30], {"Q": '"Y,N"'})
    grey_out(ws, '$Q2="Y"')

    ws = wb.create_sheet("Test times")
    table(ws, 1, MIC_TT_H, rows, set(range(5, 26)) - {11}, MIC_TT_WIDTHS, MIC_DV)
    for r in range(2, len(rows) + 2):
        ws.cell(r, 11, f'=IF(I{r}="","",IF(OR(H{r}="",J{r}=""),I{r},(H{r}+4*I{r}+J{r})/6))').number_format = "0.00"
        ws.cell(r, 19).number_format = "0%"
    link_catalog(ws, codes, MIC_LINK, "Y")
    grey_out(ws, '$E2="N"', '$U2="Y"')

    ws = wb.create_sheet("Monitoring")
    mon = [[p, a, None, c] + [None] * 9 for p, a, c in MON_ROWS] + [[None] * 13 for _ in range(15)]
    table(ws, 1, MON_H, mon, {3, 5, 6, 7, 9, 10, 11, 12, 13}, [16, 44, 12, 11, 9, 10, 10, 10, 12, 9, 11, 18, 30],
          {"D": '"' + ",".join(MIC_CATALOG) + '"', "K": '"Y,N"'})
    for r in range(2, len(mon) + 2):
        ws.cell(r, 8, f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r}*IF(G{r}="",50,G{r}))').fill = AUTO
        if r > len(MON_ROWS) + 1:
            for j in (1, 2, 4):
                ws.cell(r, j).fill = INPUT

    ws = wb.create_sheet("Support work")
    sup = [[a, u] + [None] * 6 for a, u in SUP_ROWS] + [[None] * 8 for _ in range(10)]
    table(ws, 1, SUP_H, sup, {3, 4, 5, 6, 7, 8}, [52, 12, 10, 10, 22, 10, 18, 30])
    for r in range(len(SUP_ROWS) + 2, len(sup) + 2):
        for j in (1, 2):
            ws.cell(r, j).fill = INPUT

    ws = wb.create_sheet("Extra tests")
    table(ws, 1, MIC_TT_H[:-2], [[None] * 24 for _ in range(40)], set(range(1, 25)) - {11}, MIC_TT_WIDTHS,
          dict(MIC_DV, B='"RM,FP,Water,EM,Other"'))
    for r in range(2, 42):
        ws.cell(r, 11, f'=IF(I{r}="","",IF(OR(H{r}="",J{r}=""),I{r},(H{r}+4*I{r}+J{r})/6))').number_format = "0.00"
    grey_out(ws, '$E2="N"', '$U2="Y"')

    guide(gws, wb, MIC_SHEETS, MIC_RULES, MIC_GENERAL)
    return wb, dict(rm=len(s["rm"]), n_usp=n_usp, n_nat=n_nat, fams=len(fams), products=len(s["fp"]), rows=len(rows),
                    by_code=by_code, usp=usp)


# ---------------------------------------------------------------- Markdown
def md_questions(sections):
    o = []
    for sec, qs in sections:
        o += [f"### {sec}", "", "| # | Question | Answer |", "|---|---|---|"]
        o += [f"| {q} | {t} | |" for q, t, _ in qs]
        o.append("")
    return o


USED = ("Your answers give the time each test takes and how many samples arrive, which shows how busy the lab and its "
        "instruments are and whether more staff or instruments are needed. "
        "Times are asked as minimum / most likely / maximum, and hands-on analyst time is kept apart from unattended instrument time. "
        "Do not insert, delete or reorder columns in the workbook.")


def write_md(path, title_, area, who, sheets, facts, sections, extra, used=USED):
    o = [f"# {title_}", "",
         "**Analyst edition** of the *QC Laboratory Capacity, Staffing Gap and CAPEX Assessment*.", "",
         f"- **Area:** {area}", f"- **Who fills it:** {who}",
         f"- **Working file:** `{os.path.basename(path)[:-3]}.xlsx` — the tables are only in the workbook; this page lists the questions and explains the sheets.", "",
         "## How the answers are used", "", used, "",
         "## Workbook sheets", "", "| Sheet | What to do |", "|---|---|"]
    o += [f"| {a} | {b} |" for a, b in sheets]
    o += ["", "## Pre-filled data", ""] + [f"- {f}" for f in facts] + [""] + extra
    o += ["## Questions", ""] + md_questions(sections)
    o += ["## Declaration", "", "| | |", "|---|---|", "| Respondent | |", "| Date | |", "| Reviewed by (unit head) | |", ""]
    open(path, "w", encoding="utf-8").write("\n".join(o))


def main():
    s, lead, raw = load_study(), load_leadtime(), load_raw()
    rm_wb, rm = build_rm(s, lead, raw, load_db_mapping())
    fp_wb, fp = build_fp(s)
    ph_wb, ph = build_phys(s)
    mi_wb, mi = build_micro(s)
    rm_wb.save(os.path.join(OUT, "RM_Analyst_Questionnaire.xlsx"))
    fp_wb.save(os.path.join(OUT, "FP_Analyst_Questionnaire.xlsx"))
    ph_wb.save(os.path.join(OUT, "Physical_Analyst_Questionnaire.xlsx"))
    mi_wb.save(os.path.join(OUT, "Micro_Analyst_Questionnaire.xlsx"))

    write_md(os.path.join(OUT, "RM_Analyst_Questionnaire.md"), "Raw Material Analyst Questionnaire",
             "QC Chemical Laboratory — raw materials", "each analyst testing incoming materials (2–3 people, one copy each)",
             [(a, b) for a, b, _ in RM_SHEETS if a != "Overview"],
             [f"{rm['mats']} materials on the master list ("
              + ", ".join(f"{c} {n}" for n, c in rm["cats"].most_common(4)) + ", others); "
              f"{rm['mapped']} have a USP–NF monograph.",
              f"*Test times* lists each monograph test as its own row ({rm['rows']} rows). General tests link to the "
              f"*Test catalog* ({rm['cat']} tests), so they are timed once; HPLC/GC tests are timed per material.",
              "Left out: microbial, sterility and endotoxin tests (Micro lab). Pre-set to N: tests that apply only to some grades "
              "(e.g. *where labelled for injectables*). Where a monograph gives alternative assays, keep only the one you run.",
              f"*Materials*: the master list, then {rm['sugg']} purchase-list names with a suggested match to confirm, then "
              f"{rm['other']} names found only in the purchase list or the lab time records.",
              "Recorded lab times are per test set and include unattended time; they are shown in hours as a starting point. "
              "Two columns on *Materials* flag tests in the monograph but not in the lab record, and the reverse; "
              "differences are expected where a material is tested to BP/EP or a supplier spec - then fix *Test times* to what you run."],
             RM_QUESTIONS,
             ["### Materials without a USP–NF monograph", "",
              "Tested to an in-house or supplier specification; they keep the lab's recorded test set. "
              "\\* = no recorded time either, so a fresh estimate is needed.", "",
              ", ".join(n + ("\\*" if rm["m"][n][0] == "No recorded time" else "") for n in rm["no_mono"]) + ".", ""])

    write_md(os.path.join(OUT, "FP_Analyst_Questionnaire.md"), "Finished Product Analyst Questionnaire",
             "QC Chemical Laboratory — finished products", "each analyst testing release and stability samples (2–3 people, one copy each)",
             [("Guide", "Rules, which sheets take new rows, and what every column means."),
              ("Questions", "Answer sections A–E below."),
              ("Volumes", "Release batches/yr and stability samples/yr per product, now and future; stability time points; IPC samples."),
              ("Test catalog", f"Time each general test once ({len(FP_CATALOG)} tests: appearance, water content, pH)."),
              ("Test times", f"{fp['rows']} family × test rows (assay, related substances, dissolution, CU, ID, and the general tests by dosage form)."),
              ("Campaign groups", f"{len(fp['camp'])} actives shared by more than one family: can they share an HPLC sequence?"),
              ("Extra tests", "Add tests that are missing (blend uniformity, preservative assay, delivered-dose assay).")],
             [f"{fp['products']} products in {fp['fams']} method families; dosage forms: "
              + ", ".join(f"{c} {n}" for n, c in fp["forms"].most_common()) + ".",
              f"{len(fp['no_active'])} families have no active ingredient in the product formulas ({', '.join(fp['no_active'])}); the analyst is asked to name them.",
              "Families group strengths that share methods, so times are asked once per family, not per product.",
              "Physical tests (weight, dimensions, hardness, brittleness, friability, disintegration, fill volume, IPC checks) "
              "are on the Physical lab questionnaire; microbial tests on the Micro lab questionnaire.",
              "There is no recorded time catalogue for finished products; every time on this questionnaire is a new estimate."],
             FP_QUESTIONS,
             ["### Most-shared actives (campaign candidates)", "", "| Active | Families |", "|---|---|"]
             + [f"| {a} | {f} |" for a, _, f, *_ in fp["camp"][:12]] + [""])
    write_md(os.path.join(OUT, "Physical_Analyst_Questionnaire.md"), "Physical Lab Analyst Questionnaire",
             "QC Physical Laboratory — finished products (release, stability, in-process checks)",
             "each analyst running physical tests (2–3 people, one copy each)",
             [("Guide", "Rules, which sheets take new rows, and what every column means."),
              ("Questions", "Answer sections A–E below."),
              ("Volumes", "In-process checks per batch and units per check for each product; release and stability numbers if known."),
              ("Test catalog", f"Time each physical test once ({len(PHYS_CATALOG)} tests: "
               + ", ".join(n.split(" (")[0].split(" /")[0].lower() for n, _, _ in PHYS_CATALOG.values()) + ")."),
              ("Test times", f"{ph['rows']} family × test rows, chosen by dosage form."),
              ("Extra tests", "Add tests that are missing: raw-material (particle size, bulk / tapped density), packaging "
               "(dimensions, leak test), dissolution bath work.")],
             [f"{ph['products']} products in {ph['fams']} method families; dosage forms: "
              + ", ".join(f"{c} {n}" for n, c in ph["forms"].most_common()) + ".",
              "Families group strengths that share methods, so times are asked once per family, not per product.",
              "If one tablet tester measures weight, thickness, diameter and hardness in one run, the run is timed once (on hardness).",
              "Assay, dissolution finish, water and pH are on the Chemical lab questionnaires; microbial tests on the Micro lab questionnaire.",
              "There is no recorded time catalogue for physical tests; every time on this questionnaire is a new estimate."],
             PHYS_QUESTIONS,
             ["### Physical tests by dosage form", "", "| Dosage form | Tests |", "|---|---|"]
             + [f"| {f} | " + ", ".join(PHYS_CATALOG[c][0] for c in cs) + " |" for f, cs in PHYS_FORM_TESTS.items()] + [""])

    usp_rows = [(m, ", ".join(MIC_CATALOG[c][0].split(" (")[0].replace("Absence of ", "").replace("Microbial enumeration - ", "")
                             for c, _, a, _ in t if a == "Y" and c not in ("BET", "STER"))) for m, t in mi["usp"].items()]
    write_md(os.path.join(OUT, "Micro_Analyst_Questionnaire.md"), "Micro Lab Analyst Questionnaire",
             "QC Microbiology Laboratory — raw materials, finished products, water and environmental monitoring",
             "each microbiologist (2–3 people, one copy each)",
             [("Guide", "Rules, which sheets take new rows, and what every column means."),
              ("Questions", "Answer sections A–E below."),
              ("Test catalog", f"Time each micro test once ({len(MIC_CATALOG)} tests and monitoring samples): hands-on time, "
               "incubation days, plates per sample, cabinet time."),
              ("Test times", f"{mi['rows']} raw material / product family × micro test rows: confirm Applicable and the % of lots tested."),
              ("Monitoring", "Water, environmental, personnel, compressed-air and swab programmes: points, samples per week, sampling rounds."),
              ("Support work", "Media, growth promotion, cultures, autoclave, identification, investigations, qualification: units per year and hours."),
              ("Extra tests", "Add tests that are missing (other organisms, in-house tests, stability-only tests).")],
             [f"{mi['rm']} raw materials on the master list: {mi['n_usp']} have a micro test in their USP–NF monograph; "
              f"{mi['n_nat']} more are of natural origin (extracts, premixes, gelatin) and pre-set to Y; the rest have one row set to N.",
              f"{mi['fams']} finished-product families ({mi['products']} products): TAMC / TYMC and E. coli for oral solids, plus Salmonella "
              "where the product contains a botanical; S. aureus, P. aeruginosa and preservative efficacy for the nasal spray (USP <1111>, <51>).",
              "Lots and batches per year come from the RM and FP questionnaires; this one asks only the share tested for micro.",
              "Incubation is asked in days and kept apart from analyst time: it fills incubators, not people.",
              "Sterility and endotoxin tests are pre-set to N: no sterile product is on the product list.",
              "There is no recorded time catalogue for micro tests; every time on this questionnaire is a new estimate."],
             MIC_QUESTIONS,
             ["### Raw materials with a micro test in USP–NF", "", "| Material | Tests (pre-set Y) |", "|---|---|"]
             + [f"| {m} | {t or 'none (endotoxins / sterility for parenteral grades only)'} |" for m, t in usp_rows] + [""],
             "Your answers give the hands-on time of each micro test, how long it fills incubators and cabinets, and how many "
             "samples arrive, which shows how busy the lab and its equipment are and whether more staff or equipment are needed. "
             "Times are asked as minimum / most likely / maximum, and incubation is kept apart from analyst time. "
             "Do not insert, delete or reorder columns in the workbook.")
    print(f"RM: {rm['mats']} materials, {rm['mapped']} with monograph, {rm['rows']} test rows, catalog {rm['cat']}, "
          f"suggested {rm['sugg']}, other names {rm['other']}")
    print(f"FP: {fp['products']} products, {fp['fams']} families, {fp['rows']} test rows, {len(fp['camp'])} shared actives")
    print(f"Physical: {ph['fams']} families, {ph['rows']} test rows, catalog {len(PHYS_CATALOG)}")
    print(f"Micro: {mi['n_usp']} RM with USP micro, {mi['n_nat']} natural origin, {mi['rows']} test rows, catalog {len(MIC_CATALOG)}, "
          f"RM rows set Y by code {dict(mi['by_code'])}")
    missing = [(k, n) for k, ns in ALIASES.items() for n in ns if n not in lead]
    if missing:
        print("WARNING alias targets not found:", missing)


if __name__ == "__main__":
    main()
