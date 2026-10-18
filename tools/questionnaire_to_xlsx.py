import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

src, dst = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().splitlines()

def clean(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    return t.replace("`", "").strip()

def cells(l):
    return [clean(c) for c in l.strip().strip("|").split("|")]

# blocks: (sheet title, [items]) ; item = ("h", text) | ("p", text) | ("t", header, rows, aligns)
sections, cur, i = [], None, 0
title = ""
while i < len(lines):
    l = lines[i]
    if l.startswith("# "):
        title = clean(l[2:]); cur = ("Overview", [("h", title)]); sections.append(cur)
    elif l.startswith("## "):
        h = clean(l[3:])
        m = re.match(r"([A-J])\.", h)
        name = m.group(1) + " " + re.sub(r"^[A-J]\.\s*", "", h) if m else h
        cur = (name, [("h", h)]); sections.append(cur)
    elif l.startswith("### ") or (l.startswith("**") and l.endswith("**") ) or re.match(r"\*\*[A-Z]\d\..*", l):
        cur[1].append(("s", clean(l.lstrip("# "))))
    elif l.startswith("|"):
        j = i
        while j < len(lines) and lines[j].startswith("|"): j += 1
        blk = lines[i:j]
        hdr = cells(blk[0]); aligns = [c.strip().endswith(":") for c in blk[1].strip().strip("|").split("|")]
        rows = [cells(r) for r in blk[2:]]
        cur[1].append(("t", hdr, rows, aligns)); i = j; continue
    elif l.strip() in ("", "---"):
        pass
    else:
        cur[1].append(("p", clean(re.sub(r"^- ", "• ", l))))
    i += 1

HDR = PatternFill("solid", fgColor="1F4E78"); SUB = PatternFill("solid", fgColor="DDEBF7")
INP = PatternFill("solid", fgColor="FFF9E5"); PRE = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF"); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook(); wb.remove(wb.active)
for name, items in sections:
    name = re.sub(r"[\[\]:*?/\\]", "", name)[:31]
    ws = wb.create_sheet(name)
    r, ncols = 1, 1
    widths = {}
    src_cells = []
    for it in items:
        if it[0] == "h":
            ws.cell(r, 1, it[1]).font = Font(bold=True, size=14, color="1F4E78"); r += 2
        elif it[0] == "s":
            ws.cell(r, 1, it[1]).font = Font(bold=True, size=11); r += 1
        elif it[0] == "p":
            c = ws.cell(r, 1, it[1]); c.alignment = Alignment(wrap_text=False, vertical="top"); r += 1
        else:
            _, hdr, rows, aligns = it
            n = len(hdr); ncols = max(ncols, n)
            for k, h in enumerate(hdr, 1):
                c = ws.cell(r, k, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
                c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = BRD
            ws.row_dimensions[r].height = 45; r += 1
            first = r
            for row in rows:
                row += [""] * (n - len(row))
                for k, v in enumerate(row, 1):
                    if re.fullmatch(r"-?\d+(\.\d+)?", v) and hdr[k-1] not in ("#",) and not v.startswith("0"): v = float(v) if "." in v else int(v)
                    c = ws.cell(r, k, v if v != "" else None); c.border = BRD; c.alignment = wrap
                    c.fill = PRE if v != "" else INP
                    widths[k] = max(widths.get(k, 0), min(len(str(v)), 70))
                    if aligns[k-1] and v == "": c.alignment = Alignment(horizontal="right", vertical="top")
                r += 1
            for k, h in enumerate(hdr, 1):
                widths[k] = max(widths.get(k, 0), min(len(h), 22))
                if h == "Source": src_cells.append(f"{get_column_letter(k)}{first}:{get_column_letter(k)}{r-1}")
            r += 1
    for k in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(k)].width = max(10, min(widths.get(k, 12) + 2, 70))
    if src_cells:
        dv = DataValidation(type="list", formula1='"[V],[E],[N/A],[?]"', allow_blank=True)
        dv.errorStyle = "information"; ws.add_data_validation(dv)
        for rng in src_cells: dv.add(rng)
    ws.freeze_panes = None
    ws.sheet_view.zoomScale = 100
wb.save(dst)
print([s.title for s in wb])
